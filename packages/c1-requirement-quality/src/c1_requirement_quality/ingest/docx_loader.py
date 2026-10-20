"""Word DOCX SRS loader with style hierarchy, multiline paragraphs, and table extraction."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from c1_requirement_quality.ingest.classifier import (
    MODAL_PATTERN, classify_docx_statement, is_docx_requirement_section, is_requirement_section,
)
from c1_requirement_quality.ingest.models import (
    ExtractionWarning,
    NIRDocument,
    NIRRequirement,
    NIRTable,
)

DOCX_REQ_PATTERN = re.compile(
    r"^\s*(?:[-*]\s*)?(?P<number>(?:REQ(?:UIREMENT)?[- ]?)?\d+(?:\.\d+)*|[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+)(?:[.):]|\s+-)?\s+(?P<text>[A-Za-z].*)$",
    re.IGNORECASE | re.DOTALL,
)


FUNCTIONAL_BLOCK = re.compile(r"^Functional Requirement\s+(\d+)\s*$", re.IGNORECASE)
HEADING_REQUIREMENT = re.compile(r"^(?P<id>(?:NFR|FR|REQ)[-_]\d+(?:[._-]\d+)*)(?:\s*[:.\-\u2013\u2014]\s*.+)?$", re.IGNORECASE)
ID_HEADERS = {"id", "req", "req id", "requirement id", "requirement number", "req number", "item", "item number", "number", "#"}
TEXT_HEADERS = {"requirement", "statement", "description", "specification", "text", "requirement text", "requirement statement"}


def _word_list(paragraph: Any) -> bool:
    """Recognise explicit or style-inherited Word list metadata without inventing numbering."""
    from docx.oxml.ns import qn
    style = paragraph.style
    properties = [paragraph._p.pPr]
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        properties.append(style.element.pPr)
        style = style.base_style
    return any(prop is not None and prop.find(qn("w:numPr")) is not None for prop in properties)


def load_docx_srs(path: Path) -> NIRDocument:
    """Load and parse an SRS document in Microsoft Word (.docx) format."""
    import docx

    doc = docx.Document(str(path))
    paragraphs = doc.paragraphs
    paragraph_by_element = {p._p: (i + 1, p) for i, p in enumerate(paragraphs)}
    table_by_element = {t._tbl: (i + 1, t) for i, t in enumerate(doc.tables)}
    source_blocks: list[dict[str, Any]] = []
    block_sections: dict[Any, list[str]] = {}
    headings: dict[int, str] = {}
    for element in doc.element.body:
        if element in paragraph_by_element:
            number, paragraph = paragraph_by_element[element]
            style = getattr(paragraph.style, "name", "")
            if style.startswith("Heading"):
                level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                headings = {k: v for k, v in headings.items() if k < level}
                headings[level] = paragraph.text.strip()
            source_blocks.append({"kind": "paragraph", "paragraph": number,
                                  "text": paragraph.text, "style": style, "page": None})
        elif element in table_by_element:
            number, table = table_by_element[element]
            source_blocks.append({"kind": "table", "table_id": f"Table-{number}",
                                  "rows": [[c.text for c in row.cells] for row in table.rows], "page": None})
        block_sections[element] = [headings[k] for k in sorted(headings)] or ["Unspecified"]
    normalized_paragraphs = {p._p: p.text.strip() for p in paragraphs}
    normalized_tables: dict[Any, str] = {}
    warnings: list[ExtractionWarning] = [
        ExtractionWarning(
            code="DOCX_PAGE_NUMBERS_UNAVAILABLE",
            message="DOCX format does not store physical page boundaries; requirement page numbers recorded as unknown.",
        )
    ]

    raw_text_parts: list[str] = []
    sections: list[str] = []
    current_section = "Unspecified"
    hierarchy: list[str] = ["Unspecified"]
    requirements: list[NIRRequirement] = []
    tables: list[NIRTable] = []
    context_items: list[NIRRequirement] = []
    uncertain_items: list[NIRRequirement] = []
    seen_in_section: set[str] = set()

    def append_statement(*, structured: bool = False, **fields: Any) -> None:
        category, sub_category, reason = classify_docx_statement(
            fields["raw_text"], current_section, structured=structured,
        )
        record = NIRRequirement(
            section_title=current_section, section_hierarchy=list(hierarchy),
            category=category, sub_category=sub_category, classification_reason=reason, **fields,
        )
        {"REQUIREMENT": requirements, "CONTEXT": context_items, "UNCERTAIN": uncertain_items}[category].append(record)

    # 1. Extract tables
    for tab_idx, table in enumerate(doc.tables, start=1):
        hierarchy = block_sections[table._tbl]
        current_section = hierarchy[-1]
        table_id = f"Table-{tab_idx}"
        rows_data: list[list[str]] = []
        for row in table.rows:
            rows_data.append([cell.text.strip() for cell in row.cells])

        if not rows_data:
            continue

        headers = rows_data[0]
        data_rows = rows_data[1:] if len(rows_data) > 1 else []

        tab_obj = NIRTable(
            table_id=table_id,
            headers=headers,
            rows=data_rows,
            page=None,
            line=tab_idx,
        )
        tables.append(tab_obj)
        normalized_tables[table._tbl] = tab_obj.to_markdown()

        # Inspect table rows for requirements
        id_col_idx = -1
        text_col_idx = -1
        for c_idx, h_name in enumerate(headers):
            h_lower = " ".join(h_name.lower().replace("_", " ").replace("-", " ").split())
            if h_lower in ID_HEADERS and id_col_idx == -1:
                id_col_idx = c_idx
            elif h_lower in TEXT_HEADERS and text_col_idx == -1:
                text_col_idx = c_idx

        # Schema tables are evidence/context, not a catalogue of behavioural requirements.
        schema_table = {h.strip().lower() for h in headers} >= {"attribute", "type"}
        if schema_table:
            for r_idx, cells in enumerate(data_rows, start=1):
                context_items.append(NIRRequirement(
                    doc_item_number=f"{table_id}-R{r_idx}", section_title=current_section,
                    section_hierarchy=list(hierarchy), raw_text=" | ".join(cells),
                    line=r_idx, table_id=table_id, provenance_notes=f"Row {r_idx}",
                    category="CONTEXT", sub_category="DATABASE_SCHEMA",
                    classification_reason="Database attribute/type table",
                ))
            continue

        for r_idx, r_cells in enumerate(data_rows, start=1):
            req_id_val = r_cells[id_col_idx] if 0 <= id_col_idx < len(r_cells) else ""
            req_text_val = r_cells[text_col_idx] if 0 <= text_col_idx < len(r_cells) else ""

            if not req_text_val:
                for c_val in r_cells:
                    if MODAL_PATTERN.search(c_val):
                        req_text_val = c_val
                        break

            if not req_text_val and req_id_val:
                non_id_parts = [
                    f"[{headers[c_i]}: {r_cells[c_i]}]" if c_i < len(headers) and headers[c_i] else r_cells[c_i]
                    for c_i in range(len(r_cells))
                    if c_i != id_col_idx and r_cells[c_i]
                ]
                if non_id_parts:
                    req_text_val = " ".join(non_id_parts)

            if req_text_val:
                doc_item = req_id_val if req_id_val else f"{table_id}-R{r_idx}"
                append_statement(
                    structured=bool(req_id_val or text_col_idx >= 0),
                    doc_item_number=doc_item, raw_text=req_text_val.strip(),
                    page=None, line=r_idx, table_id=table_id, provenance_notes=f"Row {r_idx}",
                )

    # 2. Extract paragraphs
    p_idx = 0
    total_paragraphs = len(paragraphs)

    while p_idx < total_paragraphs:
        p = paragraphs[p_idx]
        text = p.text.strip()
        style_name = getattr(p.style, "name", "")

        if not text:
            p_idx += 1
            continue

        # Explicit ID/title headings describe the next prose paragraph, not a
        # new governing section. Keep this narrowly scoped to requirement
        # sections and one adjacent body paragraph; never consume guidance.
        heading_id = HEADING_REQUIREMENT.fullmatch(text) if style_name.startswith("Heading") else None
        ancestors = block_sections[p._p][:-1]
        governing = ancestors[-1] if ancestors else ""
        if heading_id and is_docx_requirement_section(governing):
            cursor = p_idx + 1
            body = paragraphs[cursor] if cursor < total_paragraphs else None
            adjacent = body is not None and body._p.getprevious() is p._p
            body_text = body.text.strip() if body is not None else ""
            if (adjacent and body_text and not getattr(body.style, "name", "").startswith("Heading")
                    and not DOCX_REQ_PATTERN.match(body_text) and not _word_list(body)
                    and not body_text.endswith((":", "?"))
                    and re.search(r"\b(?:shall|must|will|should|may|can)\b", body_text, re.I)):
                current_section, hierarchy = governing, ancestors
                append_statement(
                    structured=True, doc_item_number=heading_id.group("id"), raw_text=body_text,
                    line=cursor + 1, provenance_notes=f"DOCX requirement heading paragraph {p_idx + 1}: {text}; body paragraph {cursor + 1}",
                )
                p_idx = cursor + 1
                continue
            warnings.append(ExtractionWarning(
                code="DOCX_REQUIREMENT_HEADING_WITHOUT_BODY",
                message=f"Requirement heading {heading_id.group('id')} has no supported adjacent requirement body.",
                section=governing, line=p_idx + 1,
            ))

        # Check Heading styles
        is_heading = style_name.startswith("Heading") or (
            (text.isupper() and len(text) < 70)
            or re.match(r"^\d+(?:\.\d+)*\s+[A-Z]", text) and not any(w in text.lower() for w in ("shall", "must"))
        )

        if is_heading:
            current_section = text
            if current_section not in sections:
                sections.append(current_section)
            hierarchy = block_sections[p._p]
            seen_in_section = set()
            p_idx += 1
            continue

        hierarchy = block_sections[p._p]
        current_section = hierarchy[-1]
        if HEADING_REQUIREMENT.fullmatch(current_section) and not DOCX_REQ_PATTERN.match(text):
            # A heading block owns only its explicitly associated body. Other
            # prose (including reviewer guidance) stays source context rather
            # than falling through to generic modal detection.
            context_items.append(NIRRequirement(
                doc_item_number=f"Paragraph-{p_idx + 1}", section_title=current_section,
                section_hierarchy=list(hierarchy), raw_text=text, line=p_idx + 1,
                category="CONTEXT", sub_category="EXPLANATION",
                classification_reason="Additional or unassociated prose under an explicit requirement heading",
            ))
            p_idx += 1
            continue
        block = FUNCTIONAL_BLOCK.fullmatch(text)
        if block:
            start = p_idx
            fields: dict[str, str] = {}
            cursor = p_idx + 1
            while cursor < total_paragraphs:
                candidate = paragraphs[cursor]
                if candidate._p.getprevious() is not paragraphs[cursor - 1]._p:
                    break
                field = re.match(r"^(TITLE|DESC):\s*(.+)$", candidate.text.strip(), re.IGNORECASE)
                if not field or field.group(1).upper() in fields:
                    break
                fields[field.group(1).upper()] = field.group(2)
                cursor += 1
            complete = {"TITLE", "DESC"} <= fields.keys()
            if complete:
                append_statement(
                    structured=True, doc_item_number=text, raw_text=fields["DESC"],
                    line=start + 1, provenance_notes=f"Paragraphs {start + 1}-{cursor}; {text}; TITLE: {fields['TITLE']}",
                )
                p_idx = cursor
                continue
            record = NIRRequirement(
                doc_item_number=text, section_title=current_section, section_hierarchy=list(hierarchy),
                raw_text=" ".join(q.text.strip() for q in paragraphs[start:cursor]),
                line=start + 1, provenance_notes=f"Paragraphs {start + 1}-{cursor}; {text}; TITLE: {fields.get('TITLE', '')}",
                category="UNCERTAIN", classification_reason="Incomplete functional block",
            )
            uncertain_items.append(record)
            p_idx = cursor
            continue

        section = " > ".join(hierarchy)
        relevant = is_docx_requirement_section(current_section)
        listed = _word_list(p)
        # Lists and modal-less obligations need requirement-section context.
        if (relevant or (listed and is_requirement_section(section))) and not DOCX_REQ_PATTERN.match(text):
            category, _, _ = classify_docx_statement(text, current_section)
            if category != "UNCERTAIN" or listed or MODAL_PATTERN.search(text):
                append_statement(
                    doc_item_number=f"Paragraph-{p_idx + 1}", raw_text=text, line=p_idx + 1,
                    provenance_notes="DOCX Word list paragraph" if listed else "DOCX normative paragraph",
                )
                p_idx += 1
                continue

        # Check requirement statement
        match = DOCX_REQ_PATTERN.match(text)
        if match:
            source_number = match.group("number")
            initial_text = match.group("text").strip()

            if source_number in seen_in_section:
                warnings.append(
                    ExtractionWarning(
                        code="DOCX_DUPLICATE_REQ_NUMBER",
                        message=f"Duplicate requirement number {source_number} in section '{current_section}'",
                        section=current_section,
                    )
                )
            seen_in_section.add(source_number)

            collected = [initial_text]
            category, _, _ = classify_docx_statement(initial_text, current_section, structured=True)
            req_start_p = p_idx + 1
            p_idx += 1

            # Accumulate multiline continuation paragraphs
            while category == "REQUIREMENT" and p_idx < total_paragraphs:
                cand_p = paragraphs[p_idx]
                # A table interrupts paragraph flow; never merge across it.
                if cand_p._p.getprevious() is not paragraphs[p_idx - 1]._p:
                    break
                cand_text = cand_p.text.strip()
                cand_style = getattr(cand_p.style, "name", "")

                if not cand_text:
                    p_idx += 1
                    break
                if (cand_style.startswith("Heading") or DOCX_REQ_PATTERN.match(cand_text)
                        or FUNCTIONAL_BLOCK.fullmatch(cand_text) or _word_list(cand_p)
                        or classify_docx_statement(cand_text, current_section, structured=True)[0] != "UNCERTAIN"
                        or (collected[-1].endswith((".", "!", "?")) and cand_text[0].isupper())):
                    break
                if cand_text.isupper() and len(cand_text) < 70:
                    break
                if re.match(r"^\d+(?:\.\d+)*\s+[A-Z]", cand_text) and not any(w in cand_text.lower() for w in ("shall", "must")):
                    break

                collected.append(cand_text)
                normalized_paragraphs[cand_p._p] = ""
                p_idx += 1

            # Keep the complete obligation together for lexical grounding while
            # preserving every original paragraph verbatim in source_blocks.
            normalized_paragraphs[p._p] = text + (" " + " ".join(collected[1:]) if len(collected) > 1 else "")

            append_statement(
                structured=True, doc_item_number=source_number, raw_text=" ".join(collected),
                page=None, line=req_start_p, table_id=None,
                provenance_notes=f"DOCX paragraph flow; paragraphs {req_start_p}-{p_idx}",
            )
            continue

        p_idx += 1

    for element in doc.element.body:
        text = normalized_paragraphs.get(element, normalized_tables.get(element, ""))
        if text:
            raw_text_parts.append(text)
    return NIRDocument(
        source_path=path,
        format="docx",
        raw_text="\n\n".join(raw_text_parts),
        sections=sections if sections else [current_section],
        requirements=requirements,
        tables=tables,
        context_items=context_items,
        uncertain_items=uncertain_items,
        warnings=warnings,
        source_blocks=source_blocks,
    )
