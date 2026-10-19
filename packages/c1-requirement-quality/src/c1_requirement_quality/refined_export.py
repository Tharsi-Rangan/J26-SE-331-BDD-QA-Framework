"""Readable exports from authenticated snapshots; never grant readiness."""
from __future__ import annotations

import json
import tempfile
from collections import Counter
from pathlib import Path

from c1_requirement_quality.source_integrity import sha256


def export_refined_srs(session, out: Path):
    """Create a fresh DOCX/report bundle. Unpublished drafts remain original text."""
    from docx import Document
    from c1_requirement_quality.demo_workflow import DemoSession
    session = DemoSession(session.workspace, analyser_factory=getattr(session, "_analyser_factory", None))
    if out.exists():
        raise FileExistsError("Export folder already exists; choose a new folder")
    data, _, manifest, state = session.read()
    draft = session.review_read()[3]
    initial_rag = []
    if manifest.get("rag_policy"):
        from c1_requirement_quality.workflow_persistence import MAX_SNAPSHOT_ANCESTORS, _reject_ancestry_reparse
        root, ledger = session.current, state
        for _ in range(MAX_SNAPSHOT_ANCESTORS):
            parent = ledger.get("parent")
            if parent is None:
                break
            root = root.parent / parent["artifact_path"]
            _reject_ancestry_reparse(root)
            workflow = root.parent / "internal/workflow_snapshot.json"
            _reject_ancestry_reparse(workflow)
            raw = workflow.read_bytes()
            if sha256(raw) != parent["workflow_sha256"]:
                raise ValueError("Report ancestor workflow digest mismatch")
            ledger = json.loads(raw)
        else:
            raise ValueError("Report ancestry exceeds verification limit")
        path = root.parent / "rag_diagnostics.json"
        _reject_ancestry_reparse(path)
        raw = path.read_bytes()
        if sha256(raw) != manifest["rag_diagnostics_sha256"]:
            raise ValueError("Report original RAG evidence digest mismatch")
        initial_rag = json.loads(raw)
    initial_by_id = {item["requirement_id"]: item for item in initial_rag}
    published_intents = [event for event in session.events if event["kind"] == "approval_intent"
        and event["sha256"] in state.get("rag_publication", {}).get("approval_intent_sha256s", [])]
    rows = []
    before = Counter()
    after = Counter()
    for row in data["requirements"]:
        rid = row["requirement_id"]
        record = state["requirements"][rid]
        baseline = record["original_assessment"]
        before.update(d["type"] for d in baseline.get("defects") or [])
        approved = row["human_decision"] in {"approved", "edited", "not_required"} and row["validation"]["status"] == "passed"
        display_defects = row["defects"] if approved else baseline.get("defects") or []
        after.update(d["type"] for d in display_defects)
        # The snapshot loader has independently verified any initial-RAG approval.
        final = row["final_text"] if approved else row["original_text"]
        pending = draft["requirements"][rid]
        selected = pending["candidates"].get(pending["selected_candidate_id"])
        status = "approved" if approved else "rejected" if row["human_decision"] == "rejected" or selected and selected["status"] == "REJECTED" else "unresolved"
        rows.append({"requirement_id": rid, "source_location": row.get("source"),
            "original_text": row["original_text"], "final_text": final,
            "status": status, "human_decision": row["human_decision"],
            "tier": row["tier"] if approved else baseline.get("tier", row["tier"]),
            "defects": display_defects, "original_defects": baseline["defects"],
            "candidate_review": copy_candidates(pending["candidates"]),
            "grounding": row["grounding"] if approved else baseline.get("grounding") or row["grounding"],
            "consistency": row["consistency"] if approved else baseline.get("consistency") or row["consistency"],
            "validation": row["validation"],
            "original_rag_evidence": initial_by_id.get(rid),
            "bdd_ready": bool(approved and row["bdd_ready"]), "selected_candidate_id": row.get("selected_candidate_id"),
            "clarifications": {"questions": list(pending["questions"].values()), "answers": pending["answers"],
                               "basis": "session_review_not_source_evidence"},
            "approval_intent_sha256s": [event["sha256"] for event in published_intents
                                       if event["action"]["requirement_id"] == rid]})
    counts = Counter(r["status"] for r in rows)
    report = {"report_version": 1, "source": {"filename": manifest["original_filename"],
        "document_id": manifest["document_id"], "sha256": manifest["sha256"]},
        "artifact": {"path": str(session.current), "sha256": manifest["artifact_sha256"],
                     "snapshot_id": state["snapshot_id"], "workflow_sha256": manifest.get("workflow_sha256")},
        "scope": "readable inspection export; readiness derives only from the authenticated C1 snapshot",
        "summary": {"total_requirements": len(rows), "approved": counts["approved"], "rejected": counts["rejected"],
            "unresolved": counts["unresolved"], "bdd_ready": sum(r["bdd_ready"] for r in rows),
            "quality_issues_before": sum(before.values()), "quality_issues_after": sum(after.values()),
            "smell_categories_before": dict(before), "smell_categories_after": dict(after),
            "tiers": dict(Counter(r["tier"] for r in rows)),
            "grounding_outcomes": dict(Counter(r["grounding"]["status"] for r in rows)),
            "contradictions": sum(r["consistency"]["status"] == "conflict" or r["tier"] == "TIER_3_ESCALATE" for r in rows)},
        "requirements": rows}
    report["initial_rag_diagnostics"] = initial_rag
    intent_ids = state.get("rag_publication", {}).get("approval_intent_sha256s", [])
    report["published_candidate_checks"] = [e["eligibility"] for e in session.events
        if e["kind"] == "approval_intent" and e["sha256"] in intent_ids]
    if session.meta.get("mandatory_ml"):
        report["ml_diagnostics"] = session.initial_ml_predictions()
    out.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=".c1-export-pending-", dir=out.parent))
    document = Document()
    document.add_heading("Refined Software Requirements Specification", 0)
    document.add_paragraph("Source: " + manifest["original_filename"])
    document.add_paragraph("Source SHA-256: " + manifest["sha256"])
    document.add_paragraph("Snapshot: " + state["snapshot_id"])
    summary = report["summary"]
    lines = ["# C1 Quality and Traceability Report", "", "Source: " + manifest["original_filename"],
             "Snapshot: " + state["snapshot_id"], "", json.dumps(summary, indent=2, ensure_ascii=False), ""]
    document.add_heading("Quality summary", 1)
    document.add_paragraph(f"Requirements: {len(rows)}; approved: {counts['approved']}; rejected: {counts['rejected']}; unresolved: {counts['unresolved']}; BDD-ready: {summary['bdd_ready']}.")
    document.add_paragraph(f"Quality findings: {summary['quality_issues_before']} before, {summary['quality_issues_after']} after. Unsupported details are not filled in.")
    for row in rows:
        document.add_heading(row["requirement_id"] + " — " + row["status"].upper(), 1)
        document.add_paragraph(row["final_text"])
        document.add_paragraph("Original: " + row["original_text"])
        document.add_paragraph(f"Tier: {row['tier']}; approval: {row['human_decision']}; BDD-ready: {row['bdd_ready']}.")
        document.add_paragraph("Grounding: " + row["grounding"]["status"] + "; consistency: " + row["consistency"]["status"])
        outcomes = Counter(item.get("outcome", "unresolved") for item in (row["original_rag_evidence"] or {}).get("diagnostics", []))
        if outcomes:
            document.add_paragraph("Original SRS grounding outcomes: " + str(dict(outcomes)))
        for question in row["clarifications"]["questions"]:
            document.add_paragraph("Clarification (" + question["status"] + "): " + question["text"])
        for answer in row["clarifications"]["answers"]:
            document.add_paragraph("Human answer (" + answer["actor_id"] + "): " + answer["answer"])
        lines.extend(["## " + row["requirement_id"] + " — " + row["status"].upper(),
            "Original: " + row["original_text"], "Final: " + row["final_text"],
            "Tier: " + row["tier"], "BDD-ready: " + str(row["bdd_ready"]),
            "Grounding: " + json.dumps(row["grounding"], ensure_ascii=False),
            "Original RAG evidence: " + json.dumps(row["original_rag_evidence"], ensure_ascii=False),
            "Consistency: " + json.dumps(row["consistency"], ensure_ascii=False),
            "Quality issues: " + json.dumps(row["defects"], ensure_ascii=False),
            "Clarifications: " + json.dumps(row["clarifications"], ensure_ascii=False), ""])
    document.save(stage / "refined_srs.docx")
    (stage / "quality_report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")
    (stage / "quality_report.md").write_text("\n".join(lines), encoding="utf-8")
    latest, _, current_manifest, current_state = session.read()
    if latest != data or current_manifest != manifest or current_state != state or session.review_read()[3] != draft:
        raise ValueError("Snapshot or review changed during export; staged files retained")
    if out.exists():
        raise FileExistsError(out)
    stage.rename(out)
    return {name: out / name for name in ("refined_srs.docx", "quality_report.json", "quality_report.md")}


def copy_candidates(candidates):
    return [{"candidate_id": cid, "text": item["text"], "review_status": item["status"]}
            for cid, item in candidates.items()]
