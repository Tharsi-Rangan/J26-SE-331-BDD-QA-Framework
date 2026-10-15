"""Phase 1B.1 identity, evidence integrity and publication regressions."""
from __future__ import annotations

import hashlib
import json
import runpy
from pathlib import Path

import docx
import pytest

from c1_requirement_quality import cli
from c1_requirement_quality.ingest import load_srs_document
from c1_requirement_quality.requirement_registry import RequirementRegistry
from c1_requirement_quality.source_integrity import verify_source

RUN_ID = "RUN-20261013-140000-ab12"
TEXT = "The system shall send a confirmation email."


def run_document(tmp_path, source=None, out_name="run", document_id=None):
    source = source or tmp_path / "spec.md"
    if not source.exists():
        source.write_text(f"# Requirements\n1. {TEXT}\n", encoding="utf-8")
    out = tmp_path / out_name
    cli.run(source, out, RUN_ID, registry_path=tmp_path / "registry.json", document_id=document_id)
    return out / "c1" / "validated_requirements.json", source


def test_duplicate_in_new_section_preserves_all_existing_ids(tmp_path):
    path = tmp_path / "registry.json"
    r = RequirementRegistry(path, "EX")
    first = r.assign_ids(["A", "B"], ["First", "First"], document_id="SRS")
    after = RequirementRegistry(path, "EX").assign_ids(
        ["A", "B", "A"], ["First", "First", "Second"], document_id="SRS")
    assert after[:2] == first
    assert after[2] not in first
    assert not RequirementRegistry(path, "EX").retired


def test_insert_reorder_and_renumber_preserve_cli_identity(tmp_path):
    source = tmp_path / "spec.md"
    source.write_text("# One\n1. Alpha shall return results.\n2. Beta shall return results.\n", encoding="utf-8")
    first, _ = run_document(tmp_path, source)
    before = json.loads(first.read_text())["requirements"]
    source.write_text("# One\n1. Gamma shall return results.\n2. Beta shall return results.\n3. Alpha shall return results.\n# Two\n1. Alpha shall return results.\n", encoding="utf-8")
    second, _ = run_document(tmp_path, source, "next")
    after = json.loads(second.read_text())["requirements"]
    assert after[1]["requirement_id"] == before[1]["requirement_id"]
    assert after[2]["requirement_id"] == before[0]["requirement_id"]
    assert after[3]["requirement_id"] != before[0]["requirement_id"]


def test_separate_documents_do_not_retire_each_others_ids(tmp_path):
    r = RequirementRegistry(tmp_path / "registry.json", "EX")
    a = r.assign_ids(["same"], ["S"], document_id="a")
    b = r.assign_ids(["same"], ["S"], document_id="b")
    assert a != b
    assert r.assign_ids(["new", "same"], ["S", "S"], document_id="b")[1:] == b
    assert r.documents["a"]["active"] == {'["S", "same"]': a[0]}
    assert not r.retired


def test_same_filename_at_different_locations_has_distinct_document_identity(tmp_path):
    a = tmp_path / "a" / "spec.md"; a.parent.mkdir()
    b = tmp_path / "b" / "spec.md"; b.parent.mkdir()
    for source in (a, b): source.write_text(f"1. {TEXT}\n", encoding="utf-8")
    first, _ = run_document(tmp_path, a)
    second, _ = run_document(tmp_path, b, "other")
    assert json.loads(first.read_text())["requirements"][0]["requirement_id"] != json.loads(second.read_text())["requirements"][0]["requirement_id"]


def test_explicit_document_identity_survives_source_move(tmp_path):
    first, original = run_document(tmp_path, document_id="SPEC-A")
    moved = tmp_path / "renamed.md"; moved.write_bytes(original.read_bytes())
    second, _ = run_document(tmp_path, moved, "moved", "SPEC-A")
    assert json.loads(first.read_text())["requirements"][0]["requirement_id"] == json.loads(second.read_text())["requirements"][0]["requirement_id"]


def test_indistinguishable_duplicates_fail_without_advancing_registry(tmp_path):
    path = tmp_path / "registry.json"; r = RequirementRegistry(path, "EX")
    r.assign_ids(["A"], ["S"])
    before = path.read_bytes()
    with pytest.raises(ValueError, match="indistinguishable"):
        r.assign_ids(["A", "A"], ["S", "S"])
    assert path.read_bytes() == before


def test_legacy_is_read_only_and_explicit_export_preserves_history(tmp_path):
    path = tmp_path / "legacy.json"
    legacy = {"project_code": "EX", "next_number": 4,
              "active": {"S::A": "REQ-EX-003"}, "retired": ["REQ-EX-001", "REQ-EX-002"]}
    path.write_text(json.dumps(legacy), encoding="utf-8"); before = path.read_bytes()
    r = RequirementRegistry(path, "EX")
    with pytest.raises(ValueError, match="read-only"):
        r.assign_ids(["A"], ["S"], document_id="SPEC")
    migrated = r.export_legacy(tmp_path / "new.json", "SPEC", {"S::A": ("S", "A")})
    assert migrated.assign_ids(["A", "B"], ["S", "S"], document_id="SPEC") == ["REQ-EX-003", "REQ-EX-004"]
    assert migrated.retired == {"REQ-EX-001", "REQ-EX-002"}
    assert path.read_bytes() == before
    with pytest.raises(FileExistsError): r.export_legacy(path, "SPEC", {"S::A": ("S", "A")})


def test_run_requires_explicit_registry_and_never_implicitly_migrates(tmp_path):
    srs = tmp_path / "srs.md"; srs.write_text(f"1. {TEXT}\n")
    with pytest.raises(ValueError, match="explicitly"): cli.run(srs, tmp_path / "out", RUN_ID)
    legacy = tmp_path / "legacy.json"
    legacy.write_text(json.dumps({"project_code": "EX", "next_number": 1, "active": {}, "retired": []}))
    before = legacy.read_bytes()
    with pytest.raises(ValueError, match="read-only"): cli.run(srs, tmp_path / "out", RUN_ID, registry_path=legacy)
    assert before == legacy.read_bytes()
    assert not (tmp_path / "out").exists()


def test_hash_and_full_normalized_representation_preserved(tmp_path):
    artifact, source = run_document(tmp_path)
    internal = artifact.parent / "internal"
    manifest = json.loads((internal / "source_manifest.json").read_text())
    archived = internal / manifest["archived_source"]
    assert archived.read_bytes() == source.read_bytes()
    assert manifest["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    nir = json.loads((internal / "normalized_ingestion.json").read_text())
    assert nir["requirements"][0]["raw_text"] == TEXT
    assert nir["raw_text"] == source.read_text()
    assert all(key in nir for key in ("context_items", "uncertain_items", "tables", "warnings"))
    assert verify_source(artifact, source)[0] == source.read_text()
    contract = json.loads(artifact.read_text()); cli._validate_contract_artifact(contract)
    assert (artifact.parent.parent / contract["source_document"]["path"]).read_bytes() == source.read_bytes()
    assert "sha256" not in contract["source_document"]


@pytest.mark.parametrize("command", ["refine", "generate-candidates"])
def test_changed_same_name_source_rejected_before_outputs(tmp_path, command):
    artifact, source = run_document(tmp_path)
    source.write_text("1. The system shall delete all accounts.\n")
    actions = tmp_path / "actions.json"; actions.write_text('{"requirements": []}')
    target = tmp_path / "result"
    args = [command, "--in", str(artifact), "--srs", str(source), "--out", str(target)]
    if command == "refine": args += ["--actions", str(actions)]
    assert cli.main(args) == 1
    assert not target.exists()


def test_same_bytes_renamed_source_accepted_and_bundle_survives_refinement(tmp_path):
    artifact, source = run_document(tmp_path)
    renamed = tmp_path / "copy.md"; renamed.write_bytes(source.read_bytes())
    actions = tmp_path / "actions.json"; actions.write_text('{"requirements": []}')
    out = tmp_path / "refined"
    assert cli.main(["refine", "--in", str(artifact), "--srs", str(renamed), "--actions", str(actions), "--out", str(out)]) == 0
    refined = out / "c1/validated_requirements.json"
    assert verify_source(refined, renamed)[0] == source.read_text()
    assert (out / "c1/internal/source/spec.md").read_bytes() == source.read_bytes()
    cli._validate_contract_artifact(json.loads(refined.read_text()))


@pytest.mark.parametrize("part", ["source", "normalized", "artifact"])
def test_corrupted_local_bundle_rejected(tmp_path, part):
    artifact, source = run_document(tmp_path)
    target = {"source": artifact.parent / "internal/source/spec.md",
              "normalized": artifact.parent / "internal/normalized_ingestion.json", "artifact": artifact}[part]
    target.chmod(0o644); target.write_bytes(target.read_bytes() + b" ")
    with pytest.raises(ValueError, match="hash mismatch|hash does not match"): verify_source(artifact, source)


def test_legacy_artifact_has_no_silent_filename_only_fallback(tmp_path):
    artifact, source = run_document(tmp_path)
    legacy = tmp_path / "legacy.json"; legacy.write_bytes(artifact.read_bytes())
    with pytest.raises(ValueError, match="manifest missing"): verify_source(legacy, source)


def test_docx_continuation_grounding_and_physical_blocks_preserved(tmp_path):
    source = tmp_path / "spec.docx"; d = docx.Document()
    d.add_heading("Requirements", 1)
    d.add_paragraph("1. The system shall send email")
    d.add_paragraph("within 2 seconds.")
    d.save(source)
    artifact, _ = run_document(tmp_path, source)
    requirement = json.loads(artifact.read_text())["requirements"][0]
    assert requirement["original_text"] == "The system shall send email within 2 seconds."
    assert requirement["grounding"]["status"] == "supported"
    nir = json.loads((artifact.parent / "internal/normalized_ingestion.json").read_text())
    assert "1. The system shall send email within 2 seconds." in nir["raw_text"]
    assert nir["source_blocks"][-1]["text"] == "within 2 seconds."
    assert nir["requirements"][0]["line"] == 2
    assert nir["requirements"][0]["page"] is None


def test_docx_table_order_section_and_paragraph_boundary_preserved(tmp_path):
    source = tmp_path / "spec.docx"; d = docx.Document()
    d.add_heading("Outer", 1); d.add_heading("Inner", 2)
    d.add_paragraph("1. The system shall send email.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "ID"; t.cell(0, 1).text = "Statement"
    t.cell(1, 0).text = "REQ-1"; t.cell(1, 1).text = "The system shall record events."
    d.add_paragraph("Supporting narrative after the table."); d.save(source)
    nir = load_srs_document(source)
    table_requirement = next(r for r in nir.requirements if r.table_id)
    assert table_requirement.section_hierarchy == ["Outer", "Inner"]
    paragraph_requirement = next(r for r in nir.requirements if not r.table_id)
    assert paragraph_requirement.raw_text == "The system shall send email."
    assert nir.raw_text.index("1. The system") < nir.raw_text.index("| ID |") < nir.raw_text.index("Supporting narrative")


def test_pdf_continuation_retains_requirement_start_line_and_page(tmp_path):
    helper = runpy.run_path(str(Path(__file__).parent / "ingest/test_ingest_pdf.py"))["_make_multipage_pdf"]
    source = tmp_path / "spec.pdf"
    source.write_bytes(helper([["FUNCTIONAL REQUIREMENTS", "1. The system shall send email", "within 2 seconds.", "2. The system shall record events."], ["MORE REQUIREMENTS", "3. The system shall send alerts."]]))
    nir = load_srs_document(source)
    assert nir.requirements[0].line == 2
    assert nir.requirements[0].page == 1
    assert nir.requirements[0].raw_text == "The system shall send email within 2 seconds."
    assert nir.requirements[1].line == 4
    assert nir.requirements[2].line == 2 and nir.requirements[2].page == 2
    assert nir.source_blocks[0]["page"] == 1
    assert "within 2 seconds." in nir.source_blocks[0]["text"]


@pytest.mark.parametrize("failure", ["analysis", "validation", "registry_commit"])
def test_failed_run_preserves_existing_registry_and_retains_staging(tmp_path, monkeypatch, failure):
    _, source = run_document(tmp_path)
    registry = tmp_path / "registry.json"; before = registry.read_bytes()
    source.write_text(f"1. {TEXT}\n2. The system shall record events.\n")
    def fail(*args, **kwargs): raise OSError("injected failure")
    if failure == "analysis": monkeypatch.setattr(cli, "analyse_requirements", fail)
    elif failure == "validation": monkeypatch.setattr(cli, "_validate_contract_artifact", fail)
    else: monkeypatch.setattr(RequirementRegistry, "_commit_locked", fail)
    with pytest.raises(RuntimeError, match="staging retained"):
        cli.run(source, tmp_path / "failed", RUN_ID, registry_path=registry)
    assert registry.read_bytes() == before
    assert not (tmp_path / "failed/c1").exists()
    assert list(tmp_path.glob(".c1-run-pending-*/c1/internal/source/spec.md"))


def test_stale_concurrent_registry_writer_cannot_overwrite_new_history(tmp_path):
    path = tmp_path / "registry.json"
    first = RequirementRegistry(path, "EX"); second = RequirementRegistry(path, "EX")
    first.assign_ids(["A"], document_id="a")
    before = path.read_bytes()
    second.assign_ids(["B"], document_id="b", persist=False)
    with pytest.raises(RuntimeError, match="changed while processing"): second.commit()
    assert path.read_bytes() == before


def test_existing_run_is_never_overwritten(tmp_path):
    artifact, source = run_document(tmp_path); before = artifact.read_bytes()
    registry = tmp_path / "registry.json"; previous_registry = registry.read_bytes()
    with pytest.raises(FileExistsError): cli.run(source, tmp_path / "run", RUN_ID, registry_path=registry)
    assert artifact.read_bytes() == before and registry.read_bytes() == previous_registry


def test_failed_first_run_leaves_no_registry_state(tmp_path, monkeypatch):
    source = tmp_path / "spec.md"; source.write_text(f"1. {TEXT}\n")
    def fail(*args, **kwargs): raise ValueError("invalid contract")
    monkeypatch.setattr(cli, "_validate_contract_artifact", fail)
    with pytest.raises(RuntimeError, match="invalid contract"):
        cli.run(source, tmp_path / "run", RUN_ID, registry_path=tmp_path / "registry.json")
    assert not (tmp_path / "registry.json").exists()
    assert not (tmp_path / "run/c1").exists()


def test_source_change_during_processing_rejected_before_registry_commit(tmp_path, monkeypatch):
    source = tmp_path / "spec.md"; source.write_text(f"1. {TEXT}\n")
    original = cli.analyse_requirements
    def mutate_source(*args, **kwargs):
        source.write_text("1. Changed obligation.\n")
        return original(*args, **kwargs)
    monkeypatch.setattr(cli, "analyse_requirements", mutate_source)
    with pytest.raises(RuntimeError, match="hash mismatch"):
        cli.run(source, tmp_path / "run", RUN_ID, registry_path=tmp_path / "registry.json")
    assert not (tmp_path / "registry.json").exists()
    assert not (tmp_path / "run/c1").exists()


def test_artifact_publication_failure_does_not_advance_registry(tmp_path, monkeypatch):
    _, source = run_document(tmp_path)
    registry = tmp_path / "registry.json"; before = registry.read_bytes()
    rename = Path.rename
    target = tmp_path / "next/c1"
    def fail_publication(self, destination):
        if Path(destination) == target: raise OSError("publication failed")
        return rename(self, destination)
    monkeypatch.setattr(Path, "rename", fail_publication)
    with pytest.raises(RuntimeError, match="publication failed"):
        cli.run(source, target.parent, RUN_ID, registry_path=registry)
    assert registry.read_bytes() == before and not target.exists()


def test_cli_separate_documents_leave_first_document_active(tmp_path):
    first, source = run_document(tmp_path)
    first_id = json.loads(first.read_text())["requirements"][0]["requirement_id"]
    other = tmp_path / "other.md"; other.write_text("1. The system shall record events.\n")
    run_document(tmp_path, other, "other")
    registry = RequirementRegistry(tmp_path / "registry.json", "EX")
    assert len(registry.documents) == 2
    assert first_id not in registry.retired
    assert any(first_id in document["active"].values() for document in registry.documents.values())


def test_real_registry_replacement_failure_retains_original_bytes(tmp_path, monkeypatch):
    _, source = run_document(tmp_path)
    registry = tmp_path / "registry.json"; before = registry.read_bytes()
    replace = Path.replace
    def fail_registry_replace(self, destination):
        if Path(destination) == registry: raise OSError("registry replacement failed")
        return replace(self, destination)
    monkeypatch.setattr(Path, "replace", fail_registry_replace)
    with pytest.raises(RuntimeError, match="registry replacement failed"):
        cli.run(source, tmp_path / "next", RUN_ID, registry_path=registry)
    assert registry.read_bytes() == before
    assert not (tmp_path / "next/c1").exists()
    assert list(tmp_path.glob("registry.json.pending-*"))
