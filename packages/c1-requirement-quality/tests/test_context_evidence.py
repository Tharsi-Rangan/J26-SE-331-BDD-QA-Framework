"""Context is verified evidence for explicit human review, never a decision rule."""
import asyncio
import json
import socket
from unittest.mock import patch

import docx
import pytest

from c1_requirement_quality import cli
from c1_requirement_quality.consistency_review import digest
from c1_requirement_quality.context_evidence import validate_context_reviews
from c1_requirement_quality.demo_workflow import DemoSession
from c1_requirement_quality.source_integrity import sha256

LEFT = "A student shall be allowed to borrow a maximum of three books at a time."
RIGHT = "A student shall be allowed to borrow a maximum of five books at a time."
NOTE = "The correct maximum simultaneous loan count is disputed; LIB-007 and LIB-008 intentionally conflict."
_SYSTEM_CONNECT = socket.socket.connect


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setattr(socket.socket, "connect", lambda *a, **k: pytest.fail("Network forbidden"))


def build(tmp_path, note=NOTE, left=LEFT, right=RIGHT):
    document = docx.Document()
    document.add_heading("Requirements", 1)
    document.add_paragraph("LIB-007  " + left)
    document.add_paragraph("LIB-008  " + right)
    document.add_heading("Constraints and Unresolved Decisions", 1)
    if note is not None:
        document.add_paragraph("C-04. " + note)
    document.add_paragraph("C-05. The limit for LIB-999 and LIB-998 has not been agreed.")
    source = tmp_path / "source.docx"
    document.save(source)
    cli.run(source, tmp_path / "run", "RUN-20261020-120000-s2b1", registry_path=tmp_path / "registry.json")
    return DemoSession.create(tmp_path / "session", source=source, artifact=tmp_path / "run/c1/validated_requirements.json")


@pytest.fixture
def session(tmp_path):
    return build(tmp_path)


def request(view, decision="conflict_confirmed"):
    return dict(decision=decision, evidence_refs=list(view["evidence"]), scope="Simultaneous student book loans",
                conditions="No additional condition is stated", rationale="Both source definitions disagree; the contextual note does not choose a value.",
                reviewer="synthetic-reviewer", confirmed=True, coverage="complete")


def test_verified_context_is_selectable_and_does_not_establish_the_decision(session):
    before = session.source.read_bytes()
    view = session.consistency_pairs(gate_policy=True)[0]
    context = view["evidence"]["CONTEXT:C-04"]
    assert "CONTEXT:C-05" not in view["evidence"]
    assert context["text"] == "C-04. " + NOTE
    assert context["kind"] == "verified_source_context"
    assert context["decision_effect"] == "human_review_context_only"
    assert context["provenance"]["source_block_sha256"] == sha256(context["text"].encode("utf-8"))
    assert context["provenance"]["sha256"] == sha256(before)
    assert view["current_deterministic"]["conflict_kind"] == "policy_definition_disagreement"
    assert len(view["requirements"]) == 2
    assert session.source.read_bytes() == before


def test_context_claim_of_conflict_cannot_overrule_compatible_bounds(tmp_path):
    session = build(tmp_path, left="The system shall allow at most three attempts.",
                    right="The system shall allow at most 5 attempts.")
    view = session.consistency_pairs(gate_policy=True)[0]
    assert "CONTEXT:C-04" in view["evidence"]
    assert view["current_deterministic"]["classification"] == "consistent"
    assert all(not r["bdd_ready"] for r in session.read()[0]["requirements"])


@pytest.mark.parametrize("note", [None, "The limit for LIB-007 and LIB-0080 has not been agreed.",
    "The password policy for LIB-999 and LIB-998 has not been agreed."])
def test_missing_or_unrelated_context_does_not_change_the_policy_decision(tmp_path, note):
    session = build(tmp_path, note)
    view = session.consistency_pairs(gate_policy=True)[0]
    assert not any(k.startswith("CONTEXT:") for k in view["evidence"])
    assert view["current_deterministic"]["conflict_kind"] == "policy_definition_disagreement"
    values = request(view)
    values["evidence_refs"].append("CONTEXT:C-04")
    with pytest.raises(ValueError, match="evidence references"):
        session.adopt_consistency(view["requirement_ids"], **values)


@pytest.mark.parametrize("decision", ["not_applicable", "consistent_within_scope"])
def test_policy_conflict_cannot_be_dismissed_using_context(session, decision):
    view = session.consistency_pairs(gate_policy=True)[0]
    with pytest.raises(ValueError, match="cannot be dismissed"):
        session.adopt_consistency(view["requirement_ids"], **request(view, decision))


def test_context_review_persists_resume_and_contract_gates(session, tmp_path):
    initial = session.current.read_bytes()
    view = session.consistency_pairs(gate_policy=True)[0]
    session.adopt_consistency(view["requirement_ids"], **request(view))
    resumed = DemoSession(session.workspace)
    data, _, manifest, state = resumed.read()
    assert len(data["requirements"]) == 2
    assert all(r["tier"] == "TIER_3_ESCALATE" and not r["bdd_ready"] and r["human_decision"] == "pending" for r in data["requirements"])
    assert all(r["validation"]["checks"]["consistency"] == "fail" for r in data["requirements"])
    assert all(r["escalation_active"] for r in state["requirements"].values())
    assert resumed.consistency_pairs(gate_policy=True)[0]["history"][-1]["evidence"]["CONTEXT:C-04"]["text"] == "C-04. " + NOTE
    cli._validate_contract_artifact(data)
    assert (session.workspace.parent / "run/c1/validated_requirements.json").read_bytes() == initial
    from c1_requirement_quality.refined_export import export_refined_srs
    outputs = export_refined_srs(resumed, tmp_path / "export")
    report = json.loads(outputs["quality_report.json"].read_bytes())
    assert report["summary"]["bdd_ready"] == 0


def test_audit_context_citation_replays_without_adopting_approval(session):
    view = session.consistency_pairs()[0]
    values = request(view)
    values.pop("coverage")
    before = session.current.read_bytes()
    session.record_consistency(view["requirement_ids"], **values)
    resumed = DemoSession(session.workspace)
    assert resumed.consistency_pairs()[0]["history"][-1]["evidence"]["CONTEXT:C-04"] == view["evidence"]["CONTEXT:C-04"]
    assert resumed.current.read_bytes() == before


def test_explicit_approval_cannot_override_a_policy_conflict(session):
    rid = session.read()[0]["requirements"][0]["requirement_id"]
    before = session.current.read_bytes()
    with pytest.raises(ValueError, match="escalation resolution"):
        session.apply({"requirement_id": rid, "selected_candidate_text": LEFT, "human_decision": "approved"},
                      "synthetic-reviewer", True)
    assert session.current.read_bytes() == before
    row = DemoSession(session.workspace).read()[0]["requirements"][0]
    assert row["tier"] == "TIER_3_ESCALATE"
    assert row["validation"]["checks"]["consistency"] == "fail"
    assert row["bdd_ready"] is False


@pytest.mark.parametrize("field", ["text", "doc_item_number", "source", "provenance"])
def test_forged_context_review_citations_are_rejected_even_with_new_checksum(session, field):
    from c1_requirement_quality.effective_consistency import create_review
    view = session.consistency_pairs(gate_policy=True)[0]
    review = create_review(view, request(view))
    evidence = review["evidence"]["CONTEXT:C-04"]
    if field == "text": evidence[field] = "C-04. Five is correct."
    elif field == "doc_item_number": evidence[field] = "C-99"
    elif field == "source": evidence[field]["line"] += 1
    else: evidence[field]["source_block_sha256"] = "0" * 64
    review["review_id"] = "REVIEW-" + digest({k: v for k, v in review.items() if k != "review_id"})
    data, _, manifest, _ = session.read()
    with pytest.raises(ValueError, match="Context citation"):
        validate_context_reviews(session.current.parent, data,
            {k: manifest[k] for k in ("document_id", "sha256", "normalized_sha256")}, [review])


@pytest.mark.parametrize("field", ["raw_text", "doc_item_number", "line", "section_hierarchy", "missing", "source_blocks"])
def test_forged_normalized_context_does_not_override_the_archived_source(session, field):
    root = session.current.parent / "internal"
    path = root / "normalized_ingestion.json"
    payload = json.loads(path.read_bytes())
    record = next(c for c in payload["context_items"] if c["doc_item_number"] == "C-04")
    if field == "missing": payload["context_items"].remove(record)
    elif field == "source_blocks": payload[field][4]["text"] = "C-04. Forged note"
    elif field == "line": record[field] += 1
    elif field == "section_hierarchy": record[field] = ["Forged"]
    else: record[field] = "Forged"
    path.chmod(0o644)
    path.write_text(json.dumps(payload), encoding="utf-8")
    # Even a locally recomputed NIR/workflow checksum cannot make a forged
    # context record agree with the separately archived DOCX paragraphs.
    manifest_path = root / "source_manifest.json"
    manifest = json.loads(manifest_path.read_bytes())
    manifest["normalized_sha256"] = sha256(path.read_bytes())
    state_path = root / "workflow_snapshot.json"
    state = json.loads(state_path.read_bytes())
    state["source_identity"]["normalized_sha256"] = manifest["normalized_sha256"]
    state_path.chmod(0o644)
    state_path.write_text(json.dumps(state), encoding="utf-8")
    manifest["workflow_sha256"] = sha256(state_path.read_bytes())
    manifest_path.chmod(0o644)
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="Context identity/provenance"):
        DemoSession(session.workspace).consistency_pairs(gate_policy=True)


def test_altered_archive_cannot_supply_context(session):
    source = session.current.parent / "internal/source/source.docx"
    source.chmod(0o644)
    source.write_bytes(source.read_bytes() + b"altered")
    with pytest.raises(ValueError, match="SRS hash mismatch"):
        DemoSession(session.workspace).consistency_pairs(gate_policy=True)


def test_context_visible_in_existing_textual_pair_form(session):
    from types import SimpleNamespace
    from textual.widgets import SelectionList
    from c1_requirement_quality.tui import C1App, ConsistencyScreen, ReviewForm

    before = session.current.read_bytes()

    async def journey():
        app = C1App(session=session, root=session.workspace.parent)
        async with app.run_test(size=(120, 46)) as pilot:
            owner = SimpleNamespace(session=session)
            app.push_screen(ConsistencyScreen(owner))
            await pilot.pause()
            await pilot.press("enter")
            await pilot.pause()
            assert isinstance(app.screen, ReviewForm)
            evidence = app.screen.query_one("#evidence", SelectionList)
            values = [evidence.get_option_at_index(i).value for i in range(evidence.option_count)]
            assert "CONTEXT:C-04" in values
            assert "CONTEXT:C-05" not in values
            await pilot.press("escape", "escape")

    # Match the existing Textual tests: construct Windows' internal loopback
    # socketpair before running any application code under the network block.
    with patch.object(socket.socket, "connect", _SYSTEM_CONNECT):
        loop = asyncio.new_event_loop()
    asyncio.run(journey(), loop_factory=lambda: loop)
    assert session.current.read_bytes() == before
