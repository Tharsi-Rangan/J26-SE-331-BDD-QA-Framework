"""Unit tests for Word DOCX SRS ingestion: headings, paragraph flow, and tables."""

from pathlib import Path
import docx
import pytest

from c1_requirement_quality.ingest.docx_loader import load_docx_srs
from c1_requirement_quality.srs_parser import parse_srs


def test_explicit_heading_body_blocks_preserve_ids_text_and_boundaries(tmp_path):
    doc = docx.Document()
    doc.add_heading("Functional Requirements", level=1)
    bodies = {}
    for number in range(1, 11):
        rid = f"FR-{number:03d}"
        text = ("A student may cancel a reservation within [CANCELLATION_WINDOW] hours after confirmation."
                if number == 7 else f"The system shall display reservation item {number}.")
        bodies[rid] = text
        doc.add_heading(rid + " — Reservation", level=2)
        doc.add_paragraph(text)
        doc.add_paragraph("Explanation: this paragraph must not become part of the requirement.")
    doc.add_heading("Non-Functional Requirements", level=1)
    for number in (1, 2):
        rid = f"NFR-{number:03d}"
        bodies[rid] = f"The system shall display results within {number} seconds."
        doc.add_heading(rid + " — Response time", level=2)
        doc.add_paragraph(bodies[rid])
    doc.add_heading("Questions for Stakeholder Review", level=1)
    doc.add_heading("Q-01", level=2)
    doc.add_paragraph("What should replace [CANCELLATION_WINDOW]?")
    doc.add_heading("FR-099 — Reviewer example", level=2)
    doc.add_paragraph("The system shall not treat this stakeholder example as a requirement.")
    path = _save(doc, tmp_path)
    nir = load_docx_srs(path)
    assert {r.doc_item_number: r.raw_text for r in nir.requirements} == bodies
    assert len(nir.requirements) == 12
    assert nir.requirements[6].raw_text == bodies["FR-007"]
    assert all("Explanation" not in r.raw_text for r in nir.requirements)
    assert nir.requirements[0].section_hierarchy == ["Functional Requirements"]
    assert nir.requirements[-1].section_hierarchy == ["Non-Functional Requirements"]
    assert all("heading paragraph" in r.provenance_notes and "body paragraph" in r.provenance_notes for r in nir.requirements)
    assert any(b.get("text", "").startswith("FR-001") for b in nir.source_blocks)


@pytest.mark.parametrize("prefix", ["FR", "NFR", "REQ"])
def test_heading_block_does_not_swallow_heading_guidance_or_table(tmp_path, prefix):
    doc = docx.Document()
    doc.add_heading("Requirements", level=1)
    doc.add_heading(prefix + "-001 — Missing body", level=2)
    doc.add_heading(prefix + "-002 — Actual requirement", level=2)
    doc.add_paragraph("The system shall store records.")
    doc.add_paragraph("A reviewer should check this explanation separately.")
    doc.add_heading(prefix + "-003 — Question", level=2)
    doc.add_paragraph("Which records should be retained?")
    doc.add_heading(prefix + "-004 — Interrupted by a table", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    doc.add_paragraph("The system shall not merge across a table.")
    nir = load_docx_srs(_save(doc, tmp_path))
    assert [(r.doc_item_number, r.raw_text) for r in nir.requirements] == [(prefix + "-002", "The system shall store records.")]
    assert sum(w.code == "DOCX_REQUIREMENT_HEADING_WITHOUT_BODY" for w in nir.warnings) == 3


def test_heading_body_source_ids_and_canonical_ids_stable_across_runs(tmp_path):
    import json
    from c1_requirement_quality import cli
    doc = docx.Document()
    doc.add_heading("Functional Requirements", level=1)
    doc.add_heading("REQ-001 — Store", level=2)
    doc.add_paragraph("The system shall store records.")
    source = _save(doc, tmp_path)
    before = source.read_bytes()
    results = []
    for number in (1, 2):
        out = tmp_path / f"run-{number}"
        cli.run(source, out, f"RUN-20261016-12000{number}-test", registry_path=tmp_path / "registry.json")
        results.append(json.loads((out / "c1/parsed_requirements.json").read_text()))
    assert results[0] == results[1]
    assert results[0][0]["doc_item_number"] == "REQ-001"
    assert results[0][0]["requirement_id"] == "REQ-EX-001"
    assert source.read_bytes() == before


def test_local_campus_extracts_twelve_and_kinmail_preserves_twenty_one():
    campus = Path("C:/Users/HP/Downloads/C1_Demo_SRS_Campus_Library.docx")
    kinmail = Path(__file__).resolve().parents[2] / "local-inference/kinmail-portable-001/run/c1/internal/source/SRS.docx"
    if not campus.is_file() or not kinmail.is_file():
        pytest.skip("Local reference documents unavailable; portable block tests still apply")
    before = {p: p.read_bytes() for p in (campus, kinmail)}
    result = load_docx_srs(campus)
    assert len(result.requirements) == 12
    assert [r.doc_item_number for r in result.requirements] == [f"FR-{i:03d}" for i in range(1, 11)] + ["NFR-001", "NFR-002"]
    assert "[CANCELLATION_WINDOW]" in result.requirements[6].raw_text
    assert len(load_docx_srs(kinmail).requirements) == 21
    assert all(p.read_bytes() == raw for p, raw in before.items())


def test_docx_headings_paragraphs_and_table_extraction(tmp_path: Path):
    docx_path = tmp_path / "system_spec.docx"
    doc = docx.Document()

    # Add Heading 1
    doc.add_heading("1.0 Power Subsystem", level=1)
    doc.add_paragraph("1.1 The power manager shall regulate bus voltage to 28V +/- 1V.")

    # Add Heading 2
    doc.add_heading("1.1 Battery Backup", level=2)
    doc.add_paragraph("1.2 The battery backup shall provide emergency power for at least 45 minutes.")

    # Add Table
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Requirement ID"
    table.rows[0].cells[1].text = "Statement"
    table.rows[0].cells[2].text = "Criticality"

    table.rows[1].cells[0].text = "PWR-TAB-01"
    table.rows[1].cells[1].text = "The solar array shall recharge batteries within 120 minutes."
    table.rows[1].cells[2].text = "High"

    table.rows[2].cells[0].text = "PWR-TAB-02"
    table.rows[2].cells[1].text = "The thermal controller shall dissipate excess heat above 65C."
    table.rows[2].cells[2].text = "Moderate"

    doc.save(str(docx_path))

    nir = load_docx_srs(docx_path)
    assert nir.format == "docx"
    assert len(nir.tables) == 1
    assert len(nir.requirements) == 4

    # Verify paragraph requirements
    p_reqs = [r for r in nir.requirements if r.table_id is None]
    assert len(p_reqs) == 2
    assert p_reqs[0].doc_item_number == "1.1"
    assert "regulate bus voltage" in p_reqs[0].raw_text
    assert p_reqs[1].doc_item_number == "1.2"
    assert "emergency power" in p_reqs[1].raw_text

    # Verify table requirements
    t_reqs = [r for r in nir.requirements if r.table_id is not None]
    assert len(t_reqs) == 2
    assert t_reqs[0].doc_item_number == "PWR-TAB-01"
    assert "solar array shall recharge" in t_reqs[0].raw_text
    assert t_reqs[1].doc_item_number == "PWR-TAB-02"

    # Verify warning recorded for missing page boundaries in DOCX
    assert any(w.code == "DOCX_PAGE_NUMBERS_UNAVAILABLE" for w in nir.warnings)

    # Verify compatibility parse_srs assigns canonical REQ-EX IDs
    parsed = parse_srs(docx_path, "EX")
    assert len(parsed) == 4
    for idx, item in enumerate(parsed, start=1):
        assert item["requirement_id"] == f"REQ-EX-{idx:03d}"


def _save(doc, tmp_path):
    path = tmp_path / "input.docx"
    doc.save(path)
    return path


def _list_paragraph(doc, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    paragraph = doc.add_paragraph(text)
    props = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), "1")
    numbering.append(number)
    props.append(numbering)
    return paragraph


def test_database_required_column_is_context_not_requirement_id(tmp_path):
    doc = docx.Document()
    doc.add_heading("Database Requirements", 1)
    for _ in range(2):
        table = doc.add_table(rows=2, cols=3)
        for cell, text in zip(table.rows[0].cells, ["Attribute", "Type", "Required"]):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ["id", "int", "True"]):
            cell.text = text
    nir = load_docx_srs(_save(doc, tmp_path))
    assert nir.requirements == []
    assert len(nir.tables) == len(nir.context_items) == 2
    assert [r.table_id for r in nir.context_items] == ["Table-1", "Table-2"]
    assert all(r.category == "CONTEXT" and r.line == 1 for r in nir.context_items)
    assert "True" in nir.raw_text


def test_functional_blocks_preserve_description_hierarchy_and_boundaries(tmp_path):
    doc = docx.Document()
    doc.add_heading("Functional Requirements", 1)
    for section in ["Visitors", "Members"]:
        doc.add_heading(section, 2)
        doc.add_paragraph("Functional Requirement 1")
        doc.add_paragraph("TITLE: Login")
        doc.add_paragraph("DESC: Users can log in using their credentials.")
    nir = load_docx_srs(_save(doc, tmp_path))
    assert len(nir.requirements) == 2
    assert [r.section_hierarchy for r in nir.requirements] == [
        ["Functional Requirements", "Visitors"], ["Functional Requirements", "Members"]]
    assert all(r.raw_text == "Users can log in using their credentials." for r in nir.requirements)
    assert [r.line for r in nir.requirements] == [3, 7]
    assert all("TITLE: Login" in r.provenance_notes for r in nir.requirements)
    assert any(b.get("text") == "DESC: Users can log in using their credentials." for b in nir.source_blocks)


def test_word_lists_and_unnumbered_normative_prose_without_promoting_context(tmp_path):
    doc = docx.Document()
    doc.add_heading("Performance Requirements", 1)
    doc.add_paragraph("The system should meet the following requirements:")
    texts = ["Users need to verify their email.", "The page loads fast without delay of more than one second.",
             "Only the registered users can add products."]
    for text in texts:
        _list_paragraph(doc, text)
    _list_paragraph(doc, "A diagram illustrating deployment")
    doc.add_paragraph("The system shall retain logs.")
    doc.add_heading("References", 1)
    _list_paragraph(doc, "The system shall be described in this book.")
    nir = load_docx_srs(_save(doc, tmp_path))
    assert [r.raw_text for r in nir.requirements] == texts + ["The system shall retain logs."]
    assert [r.raw_text for r in nir.uncertain_items] == ["A diagram illustrating deployment"]
    assert nir.uncertain_items[0].category == "UNCERTAIN"


def test_incomplete_functional_block_is_uncertain_and_does_not_consume_next(tmp_path):
    doc = docx.Document()
    doc.add_heading("Functional Requirements", 1)
    doc.add_paragraph("Functional Requirement 1")
    doc.add_paragraph("TITLE: Missing description")
    doc.add_paragraph("Functional Requirement 2")
    doc.add_paragraph("TITLE: Notify")
    doc.add_paragraph("DESC: The system shall send email.")
    nir = load_docx_srs(_save(doc, tmp_path))
    assert len(nir.uncertain_items) == 1
    assert "Missing description" in nir.uncertain_items[0].raw_text
    assert [r.raw_text for r in nir.requirements] == ["The system shall send email."]


def test_statement_only_table_and_legacy_continuation(tmp_path):
    doc = docx.Document()
    doc.add_heading("Requirements", 1)
    doc.add_paragraph("1. The system shall retain logs")
    doc.add_paragraph("for 30 days.")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Statement"
    table.cell(1, 0).text = "The system should encrypt backups."
    nir = load_docx_srs(_save(doc, tmp_path))
    assert {r.raw_text for r in nir.requirements} == {
        "The system shall retain logs for 30 days.", "The system should encrypt backups."}
    assert nir.requirements[0].table_id == "Table-1"


def test_broad_specification_parent_does_not_promote_schema_or_assumptions(tmp_path):
    doc = docx.Document()
    doc.add_heading("Requirement Specification", 1)
    for section, text in [("Database Requirement", "The user table is used to store users."),
                          ("User Characteristics", "They will have normal privileges."),
                          ("Assumptions and Dependencies", "The performance of the platform is good.")]:
        doc.add_heading(section, 2)
        _list_paragraph(doc, text)
    nir = load_docx_srs(_save(doc, tmp_path))
    assert nir.requirements == []
    assert len(nir.uncertain_items) == 3


def test_docx_context_identity_survives_insertion_and_duplicates_still_fail(tmp_path):
    import json
    import pytest
    from c1_requirement_quality.requirement_registry import RequirementRegistry
    doc = docx.Document()
    doc.add_heading("Security Requirements", 1)
    doc.add_paragraph("The system shall encrypt passwords.")
    path = _save(doc, tmp_path)
    registry = RequirementRegistry(tmp_path / "registry.json", "EX")
    def assign(nir):
        return registry.assign_ids([r.raw_text for r in nir.requirements],
            [json.dumps(r.section_hierarchy, ensure_ascii=False) for r in nir.requirements],
            document_id="DOC", persist=False)
    first = assign(load_docx_srs(path))
    doc.paragraphs[0].insert_paragraph_before("Supporting introduction.")
    doc.save(path)
    assert assign(load_docx_srs(path)) == first
    doc.add_paragraph("The system shall encrypt passwords.")
    doc.save(path)
    with pytest.raises(ValueError, match="Duplicate requirement text"):
        assign(load_docx_srs(path))


def test_docx_description_grounding_survives_run_refine_and_generation(tmp_path):
    import json
    from c1_requirement_quality import cli
    from c1_requirement_quality.rewrite_generator import generate_candidates_run
    doc = docx.Document()
    doc.add_heading("Functional Requirements", 1)
    doc.add_paragraph("Functional Requirement 1")
    doc.add_paragraph("TITLE: Notify")
    text = "The system shall send email. The system shall retain logs."
    doc.add_paragraph("DESC: " + text)
    source = _save(doc, tmp_path)
    original = source.read_bytes()
    out = tmp_path / "run"
    cli.run(source, out, "RUN-20261015-200000-ab12", registry_path=tmp_path / "registry.json")
    artifact_path = out / "c1/validated_requirements.json"
    artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
    row = artifact["requirements"][0]
    assert row["grounding"]["status"] == "supported"
    nir_bytes = (out / "c1/internal/normalized_ingestion.json").read_bytes()
    nir = json.loads(nir_bytes)
    assert "DESC: " + text in nir["raw_text"]
    assert row["grounding"]["evidence"][0]["text"].startswith("DESC:")
    actions = tmp_path / "actions.json"
    actions.write_text(json.dumps({"requirements": [{"requirement_id": row["requirement_id"],
        "human_decision": "not_required"}]}), encoding="utf-8")
    refined = tmp_path / "refined"
    cli.refine_run(artifact_path, source, actions, refined)
    current = json.loads((refined / "c1/validated_requirements.json").read_text(encoding="utf-8"))
    assert current["requirements"][0]["grounding"]["status"] == "supported"
    assert (refined / "c1/internal/normalized_ingestion.json").read_bytes() == nir_bytes
    class Backend:
        def generate(self, *args, **kwargs):
            return [{"text": "The system shall send email.", "strategy": "test", "source_evidence": {}}]
    suggestions = tmp_path / "suggestions.json"
    generate_candidates_run(artifact_path, source, suggestions, backend=Backend())
    result = json.loads(suggestions.read_text(encoding="utf-8"))
    assert all(r.get("rejected_reason") != "unsupported_facts" for r in result["requirements"][0]["rejected_candidates"])
    assert source.read_bytes() == original
