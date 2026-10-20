"""DOCX candidate semantics, physical boundaries and retained source evidence."""

import hashlib
import json
from pathlib import Path

import docx
import pytest

from c1_requirement_quality.ingest.docx_loader import load_docx_srs


CONSTRAINTS = [
    "Notifications must use university email. No SMS or mobile application is included in this test scope.",
    "The acceptable search response time has not been agreed.",
    "The reservation duration has not been agreed.",
    "The correct maximum simultaneous loan count is disputed; LIB-007 and LIB-008 intentionally conflict.",
    "The password-storage statement LIB-011 is intentionally insecure and must be flagged for human/security review; it is not an acceptable security policy.",
]


def save(doc, tmp_path):
    path = tmp_path / "source.docx"
    doc.save(path)
    return path


def assert_mini(nir):
    assert [r.doc_item_number for r in nir.requirements] == [f"LIB-{i:03d}" for i in range(1, 13)] + ["C-01"]
    assert [r.doc_item_number for r in nir.context_items] == ["C-02", "C-03", "C-04", "C-05"]
    assert nir.uncertain_items == []
    assert [r.line for r in nir.context_items] == [22, 23, 24, 25]
    assert [r.raw_text for r in nir.context_items] == CONSTRAINTS[1:]
    assert all(r.classification_reason for r in nir.requirements + nir.context_items)
    assert all(r.page is None for r in nir.requirements + nir.context_items)
    assert len(nir.source_blocks) == 27
    assert all(b["text"].strip() in nir.raw_text for b in nir.source_blocks if b["text"].strip())


@pytest.fixture
def mini_source(tmp_path):
    """Portable fixture with the original statement semantics and coordinates."""
    doc = docx.Document()
    doc.add_paragraph("Mini Library Management System", style="Title")
    doc.add_paragraph("Software Requirements Specification (SRS) — Controlled C1 Test Fixture", style="Subtitle")
    doc.add_heading("1. Purpose and Scope", 1)
    doc.add_paragraph("This synthetic document exercises requirement-quality review.")
    doc.add_heading("2. Actors and Business Context", 1)
    doc.add_paragraph("University email is the approved notification channel.")
    doc.add_heading("3. Functional Requirements", 1)
    statements = [
        "The system shall allow a registered student to search the library catalogue by book title.",
        "The system shall display the book title, author, and availability status for each search result.",
        "The system shall show search results quickly.",
        "The system shall allow a student to reserve an available book for [RESERVATION_HOURS] hours.",
        "The system shall send an appropriate notification to a student when a reserved book becomes available.",
        "The system shall allow a student to borrow a book and renew the loan from the same screen.",
        "A student shall be allowed to borrow a maximum of three books at a time.",
        "A student shall be allowed to borrow a maximum of five books at a time.",
        "The system shall send reservation-availability notifications to the student’s university email address.",
        "The system shall allow librarians to mark a returned book as available in the catalogue.",
        "The system shall store student passwords in plain text so that librarians can recover them.",
        "The system shall display a clear confirmation message after a student successfully reserves a book.",
    ]
    for i, text in enumerate(statements, 1):
        doc.add_paragraph(f"LIB-{i:03d}  {text}")
    doc.add_heading("4. Constraints and Unresolved Decisions", 1)
    for i, text in enumerate(CONSTRAINTS, 1):
        doc.add_paragraph(f"C-{i:02d}. {text}")
    doc.add_heading("5. Test Fixture Notice", 1)
    doc.add_paragraph("All requirements are synthetic. Do not invent answers for C-02, C-03, or C-04. Do not approve LIB-011.")
    return save(doc, tmp_path)


def test_mini_library_classification_and_verbatim_blocks(mini_source):
    before = mini_source.read_bytes()
    nir = load_docx_srs(mini_source)
    assert_mini(nir)
    original = docx.Document(mini_source)
    assert [(b["paragraph"], b["text"]) for b in nir.source_blocks] == [
        (i, p.text) for i, p in enumerate(original.paragraphs, 1)
    ]
    assert mini_source.read_bytes() == before


def test_actual_local_mini_library():
    source = Path("C:/Users/HP/Downloads/C1_Mini_Library_Test_SRS.docx")
    if not source.is_file():
        pytest.skip("Local Mini Library reference unavailable; portable fixture still runs")
    before = source.read_bytes()
    assert_mini(load_docx_srs(source))
    assert source.read_bytes() == before


@pytest.mark.parametrize("layout", ["paragraph", "table"])
@pytest.mark.parametrize("prefix", ["C", "POLICY", "REQ"])
def test_candidate_meaning_is_independent_of_id_prefix(tmp_path, layout, prefix):
    doc = docx.Document()
    doc.add_heading("Constraints and Unresolved Decisions", 1)
    if layout == "table":
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "ID", "Statement"
    for i, text in enumerate(CONSTRAINTS + ["A possible encryption approach"], 1):
        if layout == "paragraph":
            doc.add_paragraph(f"{prefix}-{i:02d}. {text}")
        else:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = f"{prefix}-{i:02d}", text
    nir = load_docx_srs(save(doc, tmp_path))
    assert [r.doc_item_number for r in nir.requirements] == [f"{prefix}-01"]
    assert [r.doc_item_number for r in nir.context_items] == [f"{prefix}-{i:02d}" for i in range(2, 6)]
    assert [r.doc_item_number for r in nir.uncertain_items] == [f"{prefix}-06"]
    records = nir.requirements + nir.context_items + nir.uncertain_items
    assert all(r.section_hierarchy == ["Constraints and Unresolved Decisions"] for r in records)
    assert all(text in nir.raw_text for text in CONSTRAINTS)
    if layout == "table":
        assert [r.line for r in records] == list(range(1, 7))
        assert all(r.table_id == "Table-1" and f"Row {r.line}" in r.provenance_notes for r in records)
        assert nir.source_blocks[-1]["rows"][1][0] == f"{prefix}-01"


@pytest.mark.parametrize("layout", ["paragraph", "table", "heading", "functional_block"])
def test_reviewer_guidance_not_promoted_by_structure_or_modals(tmp_path, layout):
    doc = docx.Document()
    doc.add_heading("Security Requirements", 1)
    text = "A reviewer must verify the password-storage decision."
    if layout == "table":
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text, table.cell(1, 0).text = "Statement", text
    elif layout == "heading":
        doc.add_heading("REQ-01 — Review", 2)
        doc.add_paragraph(text)
    elif layout == "functional_block":
        for line in ["Functional Requirement 1", "TITLE: Review", "DESC: " + text]:
            doc.add_paragraph(line)
    else:
        doc.add_paragraph(text)
    nir = load_docx_srs(save(doc, tmp_path))
    assert nir.requirements == []
    assert len(nir.context_items) == 1
    assert nir.context_items[0].raw_text == text
    assert text in nir.raw_text


def test_nfr_constraints_and_independent_paragraphs_keep_boundaries(tmp_path):
    doc = docx.Document()
    doc.add_heading("Constraints", 1)
    paragraphs = [
        "POL-01 The system shall encrypt backups.",
        "The system shall retain logs for 30 days.",
        "C-02 The system shall display results",
        "within 2 seconds.",
        "Explanation: a reviewer must verify the response-time evidence.",
        "Backup encryption is required.",
        "REQ-03 A reviewer shall be allowed to access audit records.",
        "REQ-04 The system shall flag suspicious accounts for human review.",
        "REQ-05 The system shall retain the reviewed document.",
        "REQ-07 A reviewer shall verify student identities.",
        "REQ-06 A possible deployment option",
    ]
    for text in paragraphs:
        doc.add_paragraph(text)
    nir = load_docx_srs(save(doc, tmp_path))
    assert [r.raw_text for r in nir.requirements] == [
        paragraphs[0].split(" ", 1)[1], paragraphs[1],
        "The system shall display results within 2 seconds.", paragraphs[5],
        *[text.split(" ", 1)[1] for text in paragraphs[6:10]],
    ]
    assert [r.line for r in nir.requirements] == [2, 3, 4, 7, 8, 9, 10, 11]
    assert nir.context_items[0].line == 6
    assert nir.uncertain_items[0].doc_item_number == "REQ-06"
    assert nir.source_blocks[4]["text"] == "within 2 seconds."
    assert all(text in nir.raw_text for text in paragraphs)


@pytest.mark.parametrize("label", ["", "REQ-01 "])
def test_mixed_obligation_and_missing_decision_is_uncertain(tmp_path, label):
    doc = docx.Document()
    doc.add_heading("Constraints", 1)
    text = "The system shall retain logs; the retention limit has not been agreed."
    doc.add_paragraph(label + text)
    nir = load_docx_srs(save(doc, tmp_path))
    assert nir.requirements == nir.context_items == []
    assert len(nir.uncertain_items) == 1
    assert nir.uncertain_items[0].raw_text == text
    assert label + text in nir.raw_text


def test_word_line_break_preserves_source_label_and_text(tmp_path):
    doc = docx.Document()
    doc.add_heading("Constraints", 1)
    text = "The system shall retain logs\nfor 30 days."
    doc.add_paragraph("POL-01 " + text)
    nir = load_docx_srs(save(doc, tmp_path))
    assert [(r.doc_item_number, r.raw_text, r.line) for r in nir.requirements] == [("POL-01", text, 2)]
    assert nir.source_blocks[-1]["text"] == "POL-01 " + text


@pytest.mark.parametrize("text", ["should encrypt backups.", "Which encryption should be used?"])
def test_label_and_modal_do_not_make_fragments_or_questions_requirements(tmp_path, text):
    doc = docx.Document()
    doc.add_heading("Security Requirements", 1)
    doc.add_paragraph("REQ-01 " + text)
    nir = load_docx_srs(save(doc, tmp_path))
    assert nir.requirements == nir.context_items == []
    assert [r.doc_item_number for r in nir.uncertain_items] == ["REQ-01"]
    assert text in nir.raw_text


def test_mini_pipeline_preserves_context_evidence_ids_and_contract(mini_source, tmp_path):
    from c1_requirement_quality import cli
    from c1_requirement_quality.rag.indexer import SRSIndexer
    from c1_requirement_quality.source_integrity import verify_source

    source_bytes = mini_source.read_bytes()
    out = tmp_path / "run"
    cli.run(mini_source, out, "RUN-20261020-120000-ab12", registry_path=tmp_path / "registry.json")
    root = out / "c1"
    artifact_path = root / "validated_requirements.json"
    artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
    cli._validate_contract_artifact(artifact)
    normalized_path = root / "internal/normalized_ingestion.json"
    normalized_bytes = normalized_path.read_bytes()
    nir = json.loads(normalized_bytes)
    assert len(artifact["requirements"]) == 13
    assert [r["requirement_id"] for r in artifact["requirements"]] == [f"REQ-EX-{i:03d}" for i in range(1, 14)]
    assert [r["doc_item_number"] for r in nir["context_items"]] == ["C-02", "C-03", "C-04", "C-05"]
    assert [r["line"] for r in nir["context_items"]] == [22, 23, 24, 25]
    text, manifest = verify_source(artifact_path, mini_source)
    assert manifest["sha256"] == hashlib.sha256(source_bytes).hexdigest()
    assert (root / "internal" / manifest["archived_source"]).read_bytes() == source_bytes
    assert len(nir["source_blocks"]) == 27
    chunks = SRSIndexer(text).chunk_document()
    assert all(any(statement in c.text for c in chunks) for statement in CONSTRAINTS)
    assert len({(r["source"]["section"], r["source"]["line"]) for r in artifact["requirements"]}) == 13
    diagnostics = json.loads((root / "extraction_diagnostics.json").read_text())
    assert (diagnostics["requirements_count"], diagnostics["context_items_count"], diagnostics["uncertain_items_count"]) == (13, 4, 0)

    actions = tmp_path / "actions.json"
    actions.write_text('{"requirements": []}', encoding="utf-8")
    refined = tmp_path / "refined"
    cli.refine_run(artifact_path, mini_source, actions, refined)
    assert (refined / "c1/internal/normalized_ingestion.json").read_bytes() == normalized_bytes
    assert verify_source(refined / "c1/validated_requirements.json", mini_source)[0] == text
    assert mini_source.read_bytes() == source_bytes
