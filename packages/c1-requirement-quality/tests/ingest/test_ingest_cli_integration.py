"""Integration tests for multi-format SRS execution through C1 CLI and contract schema validation."""

import json
from pathlib import Path
import docx
import pytest

from c1_requirement_quality.cli import _validate_contract_artifact, main


def _make_multipage_pdf(pages: list[list[str]]) -> bytes:
    num_pages = len(pages)
    current_obj = 3 + num_pages * 2
    font_obj_id = current_obj

    body = [b"%PDF-1.4"]
    body.append(b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj")
    kids_str = " ".join(f"{3 + i} 0 R" for i in range(num_pages))
    body.append(f"2 0 obj << /Type /Pages /Kids [{kids_str}] /Count {num_pages} >> endobj".encode("latin1"))

    for i in range(num_pages):
        page_id = 3 + i
        content_id = 3 + num_pages + i
        body.append(
            f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R /Resources << /Font << /F1 {font_obj_id} 0 R >> >> >> endobj".encode("latin1")
        )

    for i, lines in enumerate(pages):
        content_id = 3 + num_pages + i
        stream_content = "BT /F1 12 Tf 50 750 Td "
        for l in lines:
            safe_line = l.replace("(", r"\(").replace(")", r"\)")
            stream_content += f"({safe_line}) Tj 0 -15 Td "
        stream_content += "ET"
        stream_bytes = stream_content.encode("latin1")
        body.append(
            f"{content_id} 0 obj << /Length {len(stream_bytes)} >> stream\n".encode("latin1")
            + stream_bytes
            + b"\nendstream endobj"
        )

    body.append(f"{font_obj_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj".encode("latin1"))

    offsets = []
    full = bytearray()
    for obj in body:
        offsets.append(len(full))
        full.extend(obj + b"\n")
    xref_offset = len(full)
    total_objs = len(body)
    xref_table = f"xref\n0 {total_objs}\n0000000000 65535 f \n".encode("latin1") + b"".join(
        f"{off:010d} 00000 n \n".encode("latin1") for off in offsets[1:]
    )
    trailer = f"trailer << /Root 1 0 R /Size {total_objs} >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("latin1")
    return bytes(full + xref_table + trailer)


def test_cli_run_markdown_srs_contract_validation(tmp_path: Path):
    srs = tmp_path / "spec.md"
    srs.write_text(
        "# 3.0 Registration\n"
        "3.1: The system shall allow users to register.\n"
        "3.2: The system shall send confirmation emails within 5 seconds.\n",
        encoding="utf-8",
    )
    out = tmp_path / "run_md"
    ret = main([
        "run",
        "--srs", str(srs),
        "--out", str(out),
        "--run-id", "RUN-20261013-120000-ab12",
        "--registry", str(tmp_path / "reg_md.json"),
    ])
    assert ret == 0

    val_path = out / "c1" / "validated_requirements.json"
    assert val_path.is_file()
    artifact = json.loads(val_path.read_text(encoding="utf-8"))
    _validate_contract_artifact(artifact)


def test_cli_run_txt_srs_contract_validation(tmp_path: Path):
    srs = tmp_path / "spec.txt"
    srs.write_text(
        "SECTION 1: CONTROLS\n\n"
        "1.1 The flight computer shall disarm thrusters during docking.\n",
        encoding="utf-8",
    )
    out = tmp_path / "run_txt"
    ret = main([
        "run",
        "--srs", str(srs),
        "--out", str(out),
        "--run-id", "RUN-20261013-120001-ab12",
        "--registry", str(tmp_path / "reg_txt.json"),
    ])
    assert ret == 0

    val_path = out / "c1" / "validated_requirements.json"
    assert val_path.is_file()
    artifact = json.loads(val_path.read_text(encoding="utf-8"))
    _validate_contract_artifact(artifact)


def test_cli_run_json_srs_contract_validation(tmp_path: Path):
    srs = tmp_path / "spec.json"
    srs.write_text(
        json.dumps([
            {
                "id": "SYS-01",
                "section": "System Health",
                "text": "The watchdog timer shall trigger a reset if unpinned for 500 ms.",
            }
        ]),
        encoding="utf-8",
    )
    out = tmp_path / "run_json"
    ret = main([
        "run",
        "--srs", str(srs),
        "--out", str(out),
        "--run-id", "RUN-20261013-120002-ab12",
        "--registry", str(tmp_path / "reg_json.json"),
    ])
    assert ret == 0

    val_path = out / "c1" / "validated_requirements.json"
    assert val_path.is_file()
    artifact = json.loads(val_path.read_text(encoding="utf-8"))
    _validate_contract_artifact(artifact)


def test_cli_run_docx_srs_contract_validation(tmp_path: Path):
    docx_path = tmp_path / "spec.docx"
    doc = docx.Document()
    doc.add_heading("1.0 Payload Operations", level=1)
    doc.add_paragraph("1.1 The payload controller shall deploy solar panels upon separation.")
    doc.save(str(docx_path))

    out = tmp_path / "run_docx"
    ret = main([
        "run",
        "--srs", str(docx_path),
        "--out", str(out),
        "--run-id", "RUN-20261013-120003-ab12",
        "--registry", str(tmp_path / "reg_docx.json"),
    ])
    assert ret == 0

    val_path = out / "c1" / "validated_requirements.json"
    assert val_path.is_file()
    artifact = json.loads(val_path.read_text(encoding="utf-8"))
    _validate_contract_artifact(artifact)


def test_cli_run_pdf_srs_contract_validation(tmp_path: Path):
    pdf_path = tmp_path / "spec.pdf"
    pdf_path.write_bytes(
        _make_multipage_pdf([
            [
                "SECTION 1: GUIDANCE",
                "1.1 The guidance computer shall calculate trajectories at 50 Hz.",
            ]
        ])
    )

    out = tmp_path / "run_pdf"
    ret = main([
        "run",
        "--srs", str(pdf_path),
        "--out", str(out),
        "--run-id", "RUN-20261013-120004-ab12",
        "--registry", str(tmp_path / "reg_pdf.json"),
    ])
    assert ret == 0

    val_path = out / "c1" / "validated_requirements.json"
    assert val_path.is_file()
    artifact = json.loads(val_path.read_text(encoding="utf-8"))
    _validate_contract_artifact(artifact)


def test_cli_generate_candidates_with_docx_srs(tmp_path: Path):
    docx_path = tmp_path / "spec.docx"
    doc = docx.Document()
    doc.add_heading("1.0 Storage", level=1)
    doc.add_paragraph("1.1 The database shall retain audit logs for approximately 30 days.")
    doc.save(str(docx_path))

    run_dir = tmp_path / "run_initial"
    main([
        "run",
        "--srs", str(docx_path),
        "--out", str(run_dir),
        "--run-id", "RUN-20261013-120005-ab12",
        "--registry", str(tmp_path / "reg.json"),
    ])

    val_path = run_dir / "c1" / "validated_requirements.json"
    assert val_path.is_file()

    cand_out = tmp_path / "candidate_suggestions.json"
    ret = main([
        "generate-candidates",
        "--in", str(val_path),
        "--srs", str(docx_path),
        "--out", str(cand_out),
    ])
    assert ret == 0
    assert cand_out.is_file()

    cand_data = json.loads(cand_out.read_text(encoding="utf-8"))
    assert cand_data["contract"] == "c1.candidate-suggestions"
    assert cand_data["schema_version"] == "1.0.0"
    assert len(cand_data["requirements"]) == 1
