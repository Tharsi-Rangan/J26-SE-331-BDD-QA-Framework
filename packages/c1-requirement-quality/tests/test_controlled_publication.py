"""Controlled publication uses real integrity logic and injected offline evidence."""
import json
import pytest
import copy
from pathlib import Path

from test_candidate_validation import context
from test_grounded_review_intents import review, select, register, approve
from c1_requirement_quality.demo_workflow import DemoSession
from c1_requirement_quality.rag.publication import publication_models
from c1_requirement_quality.source_integrity import verify_source
from c1_requirement_quality.refined_export import export_refined_srs
from c1_requirement_quality import cli
from c1_requirement_quality.wizard import Wizard


@pytest.fixture(autouse=True)
def offline_provenance(review):
    with publication_models(review[4]):
        yield


def test_valid_publication_restart_and_readiness(review):
    session, rid, candidate, control, factory = review
    original = session.current
    select(review)
    register(review)
    approve(review)
    analyses = control['analyses']
    session.publish_approved('Publisher', True)
    assert session.current != original
    restarted = DemoSession(session.workspace, analyser_factory=factory)
    data, _, manifest, state = restarted.read()
    row = data['requirements'][0]
    assert row['final_text'] == candidate
    assert row['human_decision'] == 'approved'
    assert row['bdd_ready']
    assert state['rag_publication']['approval_intent_sha256s']
    assert manifest['rag_ancestry_version'] == 1
    assert restarted.pending is None
    assert control['analyses'] == analyses
    verify_source(restarted.current, restarted.source)
    with pytest.raises(ValueError, match='No current approved'):
        restarted.publish_approved('Publisher', True)


@pytest.mark.parametrize('state', ['none', 'selected', 'registered', 'rejected', 'stale_question', 'stale_selection'])
def test_unapproved_and_stale_publications_blocked(review, state):
    session, rid, candidate, _, _ = review
    if state != 'none':
        select(review)
    if state not in {'none', 'selected'}:
        register(review)
    if state.startswith('stale'):
        approve(review)
        if state == 'stale_question':
            session.review_action({'requirement_id': rid, 'clarification_questions': ['Which logs?']}, 'Other', True)
        else:
            select(review, candidate + '.')
    if state == 'rejected':
        session.review_action({'requirement_id': rid, 'human_decision': 'rejected'}, 'Reviewer', True)
    with pytest.raises(ValueError, match='No current approved'):
        session.publish_approved('Publisher', True)
    assert session.pending is None


@pytest.mark.parametrize('part', ['source', 'receipt', 'reference', 'intent', 'policy'])
def test_modified_evidence_cannot_publish(review, part, monkeypatch):
    session, _, _, _, _ = review
    select(review); register(review); approve(review)
    if part == 'source':
        session.source.write_bytes(session.source.read_bytes() + b' ')
    elif part in {'receipt', 'reference'}:
        event = next(e for e in session.events if e['kind'] == 'candidate_evidence_registered')
        path = session.workspace / event['reference']
        if part == 'receipt':
            path = session.workspace / json.loads(path.read_bytes())['evidence']
        path.write_bytes(path.read_bytes() + b' ')
    elif part == 'intent':
        path = session.workspace / 'events' / f'{len(session.events):06}.json'
        data = json.loads(path.read_bytes()); data['action']['selected_candidate_id'] = 'forged'
        path.write_text(json.dumps(data), encoding='utf-8')
    else:
        from c1_requirement_quality.rag import candidate_validation
        monkeypatch.setattr(candidate_validation, 'verification_policy', lambda: {'version': 'changed'})
    with pytest.raises((ValueError, OSError)):
        session.publish_approved('Publisher', True)
    assert not (session.workspace / 'snapshots').exists()


@pytest.mark.parametrize('point', ['before_prepare', 'before_rename', 'after_rename'])
def test_publication_interruption_and_recovery(review, monkeypatch, point):
    session, _, _, _, factory = review
    select(review); register(review); approve(review)
    parent = session.current
    with monkeypatch.context() as patch:
        if point == 'before_prepare':
            import c1_requirement_quality.workflow_persistence as persistence
            patch.setattr(persistence, 'archive_parent_evidence_links', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('interrupted before prepare')))
        elif point == 'before_rename':
            rename = Path.rename
            def interrupt(path, target):
                if path.name.startswith('.c1-refine-pending-'):
                    raise RuntimeError('interrupted before rename')
                return rename(path, target)
            patch.setattr(Path, 'rename', interrupt)
        else:
            patch.setattr(session, 'recover', lambda *a: (_ for _ in ()).throw(RuntimeError('interrupted after rename')))
        with pytest.raises(RuntimeError, match='interrupted'):
            session.publish_approved('Publisher', True)
    restarted = DemoSession(session.workspace, analyser_factory=factory)
    assert restarted.pending
    assert restarted.current == parent
    if point == 'after_rename':
        restarted.recover(True)
        current = restarted.current
        restarted.recover(True)
        assert DemoSession(restarted.workspace, analyser_factory=factory).current == current
        assert restarted.read()[0]['requirements'][0]['bdd_ready']
    else:
        with pytest.raises((ValueError, OSError)):
            restarted.recover(True)
        restarted.recover(False)
        assert restarted.current == parent


def test_direct_api_cannot_forge_operation_or_approval(review, tmp_path):
    session, rid, candidate, _, _ = review
    with pytest.raises(ValueError):
        session.apply({'requirement_id': rid, 'human_decision': 'approved'}, 'Attacker', True,
            _approval_intent_sha256s=['forged'], evidence_candidate_ids=['forged'])
    actions = tmp_path / 'actions.json'
    actions.write_text('{"requirements": []}', encoding='utf-8')
    with pytest.raises((ValueError, KeyError)):
        cli.refine_run(session.current, session.source, actions, tmp_path / 'forged',
            _operation={'approval_intent_sha256s': ['forged']}, _prepared=lambda c: {})
    assert not (tmp_path / 'forged').exists()
    from c1_requirement_quality.rag.publication import materialize
    with pytest.raises(ValueError, match='Journal-backed'):
        materialize({'approval_intent_sha256s': ['forged']}, session.source)


def test_cancelled_prepared_output_cannot_authorize_readiness(review, monkeypatch):
    session, _, _, _, factory = review
    select(review); register(review); approve(review)
    with monkeypatch.context() as patch:
        patch.setattr(session, 'recover', lambda *a: (_ for _ in ()).throw(RuntimeError('interrupt')))
        with pytest.raises(RuntimeError, match='interrupt'):
            session.publish_approved('Publisher', True)
    restarted = DemoSession(session.workspace, analyser_factory=factory)
    child = Path(restarted.pending['out']) / 'c1/validated_requirements.json'
    restarted.recover(False)
    with pytest.raises(ValueError, match='Cancelled publication'):
        verify_source(child, restarted.source)


def test_published_bundle_and_inspection_copy_cannot_be_substituted(review, tmp_path):
    session, _, _, _, _ = review
    select(review); register(review); approve(review)
    session.publish_approved('Publisher', True)
    import shutil
    copy_path = session.current.parent.parent.parent / 'copied/c1'
    shutil.copytree(session.current.parent, copy_path)
    with pytest.raises(ValueError, match='Inspection or copied'):
        verify_source(copy_path / session.current.name, session.source)
    path = session.current
    path.chmod(0o644)
    path.write_bytes(path.read_bytes() + b' ')
    with pytest.raises(ValueError, match='hash'):
        session.read()


def test_refined_docx_and_quality_report_match_published_snapshot(review, tmp_path):
    from docx import Document
    session, rid, candidate, _, _ = review
    select(review); register(review); approve(review)
    session.publish_approved('Publisher', True)
    outputs = export_refined_srs(session, tmp_path / 'export')
    text = '\n'.join(p.text for p in Document(outputs['refined_srs.docx']).paragraphs)
    assert rid in text and candidate in text and 'APPROVED' in text
    report = json.loads(outputs['quality_report.json'].read_bytes())
    assert report['summary']['approved'] == 1
    assert report['summary']['bdd_ready'] == 1
    assert report['requirements'][0]['final_text'] == candidate
    assert report['requirements'][0]['original_text'] != candidate
    assert report['initial_rag_diagnostics'][0]['requirement_id'] == rid
    assert report['requirements'][0]['original_rag_evidence']['diagnostics']
    assert report['published_candidate_checks'][0]['candidate_id'] == session.read()[0]['requirements'][0]['selected_candidate_id']
    assert outputs['quality_report.md'].is_file()
    with pytest.raises(FileExistsError):
        export_refined_srs(session, tmp_path / 'export')


def test_unpublished_intent_export_never_approves_or_refines(review, tmp_path):
    from docx import Document
    session, _, candidate, _, _ = review
    select(review); register(review); approve(review)
    outputs = export_refined_srs(session, tmp_path / 'inspection')
    report = json.loads(outputs['quality_report.json'].read_bytes())
    assert report['summary']['bdd_ready'] == 0
    assert report['summary']['approved'] == 0
    assert report['requirements'][0]['final_text'] == session.read()[0]['requirements'][0]['original_text']
    assert 'UNRESOLVED' in '\n'.join(p.text for p in Document(outputs['refined_srs.docx']).paragraphs)


def test_publication_and_export_cli(review, tmp_path, capsys):
    session, _, _, _, _ = review
    select(review); register(review); approve(review)
    assert cli.main(['publish-approved', '--session', str(session.workspace), '--reviewer', 'Publisher', '--confirm', 'yes']) == 0
    assert cli.main(['export-refined', '--session', str(session.workspace), '--out', str(tmp_path / 'cli-export')]) == 0
    assert (tmp_path / 'cli-export/refined_srs.docx').is_file()
    assert 'authorized C1 output' in capsys.readouterr().out


def test_mixed_approved_and_unresolved_requirements_export(review, tmp_path, monkeypatch):
    from docx import Document
    _, _, _, control, factory = review
    class Relevant(factory):
        def analyse(self, text):
            old = control['passages']
            control['passages'] = [text.replace('shall', 'must')]
            try:
                return super().analyse(text)
            finally:
                control['passages'] = old
    source = tmp_path / 'mixed.md'
    source.write_text('# Requirements\n1. The system shall accept passwords of at least 8 characters.\n'
                      '2. The system shall accept passwords of at most 64 characters.\n', encoding='utf-8')
    monkeypatch.setattr(cli, 'RAGGroundingAnalyser', Relevant)
    cli.run(source, tmp_path / 'mixed-run', 'RUN-20261019-120001-abcd', registry_path=tmp_path / 'mixed-registry.json',
            grounding_mode='rag', mandatory_rag=True)
    session = DemoSession.create(tmp_path / 'mixed-session', artifact=tmp_path / 'mixed-run/c1/validated_requirements.json', source=source)
    session = DemoSession(session.workspace, analyser_factory=Relevant)
    rows = session.read()[0]['requirements']
    assert len(rows) == 2
    rid = rows[0]['requirement_id']
    candidate = rows[0]['original_text'].rstrip('.')
    session.review_action({'requirement_id': rid, 'selected_candidate_text': candidate}, 'Reviewer', True)
    session.validate_candidate_evidence(rid, candidate)
    session.review_action({'requirement_id': rid, 'human_decision': 'approved'}, 'Reviewer', True)
    session.publish_approved('Publisher', True)
    outputs = export_refined_srs(session, tmp_path / 'mixed-export')
    report = json.loads(outputs['quality_report.json'].read_bytes())
    assert report['summary']['total_requirements'] == 2
    assert report['summary']['approved'] == 1
    assert report['summary']['unresolved'] == 1
    assert report['summary']['bdd_ready'] == 1
    assert report['requirements'][1]['final_text'] == rows[1]['original_text']
    assert report['requirements'][0]['approval_intent_sha256s']
    assert not report['requirements'][1]['approval_intent_sha256s']
    text = '\n'.join(p.text for p in Document(outputs['refined_srs.docx']).paragraphs)
    assert 'APPROVED' in text and 'UNRESOLVED' in text


def test_legacy_document_export_retains_originals(review, tmp_path):
    source = review[0].source
    cli.run(source, tmp_path / 'legacy', 'RUN-20261019-120002-abcd', registry_path=tmp_path / 'legacy-registry.json')
    session = DemoSession.create(tmp_path / 'legacy-session', artifact=tmp_path / 'legacy/c1/validated_requirements.json', source=source)
    outputs = export_refined_srs(session, tmp_path / 'legacy-export')
    report = json.loads(outputs['quality_report.json'].read_bytes())
    assert report['requirements'][0]['final_text'] == report['requirements'][0]['original_text']


def test_wizard_explicit_publish_action(review, tmp_path):
    session, rid, _, _, _ = review
    select(review); register(review); approve(review)
    values = iter(['Publisher', 'yes'])
    output = []
    ui = Wizard(tmp_path / 'ui', ask=lambda _: next(values), tell=output.append)
    ui.publish_approved(session, rid)
    assert session.read()[0]['requirements'][0]['bdd_ready']
    assert any('Published verified refinement' in line for line in output)


def test_wizard_can_review_original_without_inventing_a_rewrite(review, tmp_path):
    session, rid, _, _, _ = review
    values = iter(['1', 'Reviewer', 'yes'])
    ui = Wizard(tmp_path / 'ui', ask=lambda _: next(values), tell=lambda _: None)
    ui.approval(session, rid, original_only=True)
    assert session.review_read()[0]['requirements'][0]['selected_candidate_id']
    assert not session.approval_intents()
    values = iter(['1', 'Reviewer', 'yes'])
    ui.approval(session, rid)
    session.publish_approved('Publisher', True)
    row = session.read()[0]['requirements'][0]
    assert row['final_text'] == row['original_text'] and row['bdd_ready']


def test_second_generation_publication_reverifies_ancestry(review):
    session, rid, _, _, factory = review
    select(review); register(review); approve(review)
    session.publish_approved('Publisher', True)
    first = session.current
    original = session.read()[0]['requirements'][0]['original_text']
    session.review_action({'requirement_id': rid, 'selected_candidate_text': original}, 'Reviewer', True)
    session.validate_candidate_evidence(rid, original)
    session.review_action({'requirement_id': rid, 'human_decision': 'approved'}, 'Reviewer', True)
    session.publish_approved('Publisher', True)
    restarted = DemoSession(session.workspace, analyser_factory=factory)
    assert restarted.current != first
    assert restarted.read()[3]['round'] == 2
    assert restarted.read()[0]['requirements'][0]['bdd_ready']
