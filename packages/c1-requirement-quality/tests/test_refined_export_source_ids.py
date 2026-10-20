"""Original document labels survive export alongside canonical C1 IDs."""
import json

from docx import Document

from c1_requirement_quality import cli
from c1_requirement_quality.demo_workflow import DemoSession
from c1_requirement_quality.refined_export import export_refined_srs


def test_docx_source_ids_survive_unresolved_export(tmp_path):
    source = tmp_path / 'source.docx'
    document = Document()
    document.add_heading('Functional Requirements', 1)
    document.add_paragraph('FR-01  The system shall send email.')
    document.add_paragraph('FR-02  The system shall record bookings.')
    document.save(source)
    cli.run(source, tmp_path / 'run', 'RUN-20261020-120000-test',
            registry_path=tmp_path / 'registry.json')
    session = DemoSession.create(tmp_path / 'session',
        artifact=tmp_path / 'run/c1/validated_requirements.json')
    paths = export_refined_srs(session, tmp_path / 'export')
    report = json.loads(paths['quality_report.json'].read_text(encoding='utf-8'))
    assert [r['source_document_item'] for r in report['requirements']] == ['FR-01', 'FR-02']
    assert all(r['original_text'] == r['final_text'] and not r['bdd_ready']
               for r in report['requirements'])
    text = '\n'.join(p.text for p in Document(paths['refined_srs.docx']).paragraphs)
    markdown = paths['quality_report.md'].read_text(encoding='utf-8')
    for label in ('FR-01', 'FR-02'):
        assert 'Original document ID: ' + label in text
        assert 'Original document ID: ' + label in markdown
